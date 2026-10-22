# ============================================================
# DayBill ERP - Report Center (data layer)
# ------------------------------------------------------------
# Every business report in one place, as plain functions that return
#   {"title", "subtitle", "columns", "rows", "totals", "summary", "notes"}
# The Reports screen (ui/report_center.py), the VAT Report screen and
# the Claude connector all use these, so a figure is the same wherever
# it is shown.
#
# Rules shared by all reports
#   * only CONFIRMED documents count (drafts / cancelled are ignored)
#   * accounting reports read POSTED journal entries
#   * quantities are in each product's BASE unit (Box of 12 = 12)
#   * cost = product average cost (else purchase price)
# ============================================================

import sys, os
import datetime as _dt
from collections import OrderedDict, defaultdict

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

try:
    from config.database import db as _app_db
except Exception:                     # pragma: no cover
    _app_db = None

_DB = {"db": _app_db}
_CACHE = {}


def use_db(obj):
    """Point the reports at another db object (the Claude connector)."""
    _DB["db"] = obj
    _CACHE.clear()
    try:                                  # default-account lookups too
        from modules import account_defaults as _ad
        _ad.db = obj
        _ad.invalidate()
    except Exception:
        pass


def _db():
    return _DB["db"]


def _q(sql, params=()):
    return _db().fetchall(sql, tuple(params)) or []


def _q1(sql, params=()):
    r = _q(sql, params)
    return r[0] if r else None


def _split(h):
    """(cash, card) actually taken at the till on a POS bill (cash + card
    fields saved from v5.15), else None for older bills."""
    if h.get("cash_paid") is None and h.get("card_paid") is None:
        return None
    return (_f(h.get("cash_paid")), _f(h.get("card_paid")))


def _f(v):
    try:
        return float(v or 0)
    except Exception:
        return 0.0


def _d(v):
    if v is None or v == "":
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10]


def _has_col(table, col):
    key = ("col", table, col)
    if key not in _CACHE:
        try:
            _q(f"SELECT {col} FROM {table} WHERE 1=0")
            _CACHE[key] = True
        except Exception:
            _CACHE[key] = False
    return _CACHE[key]


def _has_table(table):
    key = ("tbl", table)
    if key not in _CACHE:
        try:
            _q(f"SELECT COUNT(*) AS n FROM {table} WHERE 1=0")
            _CACHE[key] = True
        except Exception:
            _CACHE[key] = False
    return _CACHE[key]


# ── columns / result helpers ──────────────────────────────────
NUMERIC = ("money", "qty", "pct", "int")


def col(key, label, kind="text", width=110):
    return {"key": key, "label": label, "kind": kind, "width": width,
            "align": "e" if kind in NUMERIC else "w"}


def _result(title, columns, rows, totals=None, summary=None, notes=None,
            subtitle=""):
    return {"title": title, "subtitle": subtitle, "columns": columns,
            "rows": rows, "totals": totals, "summary": summary or [],
            "notes": notes or []}


def _sum_totals(rows, keys, label_key=None, label="TOTAL"):
    t = {k: round(sum(_f(r.get(k)) for r in rows if not r.get("_style")), 2)
         for k in keys}
    if label_key:
        t[label_key] = label
    return t


def _period_text(f, t):
    return f"{f}  →  {t}"


# ── shared lookups ────────────────────────────────────────────
def _factor_map():
    fm = {}
    if not _has_table("product_units"):
        return fm
    for r in _q("SELECT product_id, unit_name_en, conversion_factor, is_active "
                "FROM product_units"):
        k = (r["product_id"], (r.get("unit_name_en") or "").strip().lower())
        f = _f(r.get("conversion_factor")) or 1.0
        if k not in fm or r.get("is_active"):
            fm[k] = f
    return fm


def _factor(fm, pid, uname):
    return fm.get((pid, (uname or "").strip().lower()), 1.0)


def _products():
    return {r["product_id"]: r for r in _q("SELECT * FROM products")}


def _cost(p):
    if not p:
        return 0.0
    for k in ("average_cost", "purchase_price", "last_purchase_cost"):
        if _f(p.get(k)) > 0:
            return _f(p.get(k))
    return 0.0


def _categories():
    try:
        return {r["category_id"]: r.get("category_name_en") or ""
                for r in _q("SELECT category_id, category_name_en FROM product_categories")}
    except Exception:
        return {}


def category_names():
    return sorted(v for v in _categories().values() if v)


def _users():
    try:
        return {r["user_id"]: (r.get("full_name_en") or r.get("username") or "")
                for r in _q("SELECT user_id, username, full_name_en FROM users")}
    except Exception:
        return {}


def accounts_list(kind="all"):
    """[(code, name)] for pickers. kind: all | cashbank."""
    try:
        rows = _q("SELECT account_code, account_name_en, account_type FROM accounts "
                  "ORDER BY account_code")
    except Exception:
        return []
    if kind == "cashbank":
        rows = [r for r in rows if r.get("account_type") == "asset"
                and str(r.get("account_code", "")).startswith("11")]
    return [(r["account_code"], r["account_name_en"]) for r in rows]


def _account(code_or_id):
    if code_or_id in (None, ""):
        return None
    s = str(code_or_id).split(" ")[0].strip()
    r = _q1("SELECT account_id, account_code, account_name_en, account_type FROM accounts "
            "WHERE account_code=?", (s,))
    if r:
        return r
    if s.isdigit():
        return _q1("SELECT account_id, account_code, account_name_en, account_type "
                   "FROM accounts WHERE account_id=?", (int(s),))
    return None


def _default_code(role):
    try:
        from modules import account_defaults as ad
        aid = getattr(ad, role)()
        r = _q1("SELECT account_code FROM accounts WHERE account_id=?", (aid,))
        return r["account_code"] if r else None
    except Exception:
        return {"cash": "1100", "bank": "1110"}.get(role)


def _match(text, *vals):
    t = (text or "").strip().lower()
    if not t:
        return True
    return any(t in str(v or "").lower() for v in vals)


# ── documents (confirmed) ─────────────────────────────────────
def _sales_headers(f, t):
    return _q("SELECT * FROM sales_invoices WHERE invoice_status='confirmed' "
              "AND invoice_date BETWEEN ? AND ? ORDER BY invoice_date, invoice_id", (f, t))


def _sales_lines(f, t):
    return _q("SELECT sil.*, si.invoice_date, si.invoice_number, si.customer_id, "
              "si.customer_name_en AS cust_name, si.invoice_type_code "
              "FROM sales_invoice_lines sil JOIN sales_invoices si ON si.invoice_id=sil.invoice_id "
              "WHERE si.invoice_status='confirmed' AND si.invoice_date BETWEEN ? AND ?", (f, t))


def _is_credit_note(h):
    return str(h.get("invoice_type_code") or "388") == "381"


def _return_headers(kind, f, t):
    tbl = "sales_returns" if kind == "sales" else "purchase_returns"
    if not _has_table(tbl):
        return []
    return _q(f"SELECT * FROM {tbl} WHERE ISNULL(return_status,'confirmed')='confirmed' "
              f"AND return_date BETWEEN ? AND ? ORDER BY return_date, return_id", (f, t))


def _return_lines(kind, f, t):
    tbl = "sales_returns" if kind == "sales" else "purchase_returns"
    lt = "sales_return_lines" if kind == "sales" else "purchase_return_lines"
    if not _has_table(tbl):
        return []
    return _q(f"SELECT l.*, r.return_date, r.return_number FROM {lt} l "
              f"JOIN {tbl} r ON r.return_id=l.return_id "
              f"WHERE ISNULL(r.return_status,'confirmed')='confirmed' "
              f"AND r.return_date BETWEEN ? AND ?", (f, t))


def _purchase_headers(f, t):
    return _q("SELECT * FROM purchase_invoices WHERE invoice_status='confirmed' "
              "AND invoice_date BETWEEN ? AND ? ORDER BY invoice_date, purchase_id", (f, t))


def _purchase_lines(f, t):
    return _q("SELECT pil.*, pi.invoice_date, pi.invoice_number, pi.supplier_id, "
              "pi.supplier_name_en AS sup_name FROM purchase_invoice_lines pil "
              "JOIN purchase_invoices pi ON pi.purchase_id=pil.purchase_id "
              "WHERE pi.invoice_status='confirmed' AND pi.invoice_date BETWEEN ? AND ?", (f, t))


# ── journals ──────────────────────────────────────────────────
TYPE_LABEL = {"sales": "Sale", "pos": "POS Sale", "cogs": "Cost of Sales",
              "purchase": "Purchase", "sales_return": "Sales Return",
              "cogs_return": "Return Cost", "purchase_return": "Purchase Return",
              "receipt": "Receipt", "payment": "Payment", "expense": "Expense",
              "adjustment": "Adjustment", "stock_adjust": "Stock Adjust",
              "transfer": "Transfer", "salary_payment": "Salary", "payroll": "Payroll",
              "manual": "Journal", "voucher": "Voucher", "opening": "Opening",
              "opening_customer": "Opening Balance", "opening_supplier": "Opening Balance",
              "opening_stock": "Opening Stock", "opening_accounts": "Opening Balance",
              "purchase_return_cost": "Return Price Difference"}


def _vt_col():
    from modules import voucher_types as _vt
    return _vt.col()


def vlabel(row):
    """CR · Cash Receipt — the voucher type of a journal row (v6.9.3)."""
    from modules import voucher_types as _vt
    return _vt.label(_vt.of_journal(row))


def type_label(st):
    return TYPE_LABEL.get(str(st or "").lower(), str(st or "Journal").replace("_", " ").title())


def _acc_movement(account_id, f, t):
    op = _q1("SELECT ISNULL(SUM(jl.debit),0) AS dr, ISNULL(SUM(jl.credit),0) AS cr "
             "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
             "WHERE je.is_posted=1 AND jl.account_id=? AND je.entry_date < ?",
             (account_id, f)) or {}
    rows = _q("SELECT je.journal_id, je.entry_date, je.journal_number, je.source_type, "
              "je.reference, je.description_en, jl.description AS line_desc, jl.debit, jl.credit "
              "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
              "WHERE je.is_posted=1 AND jl.account_id=? AND je.entry_date BETWEEN ? AND ? "
              "ORDER BY je.entry_date, je.journal_id, jl.line_id", (account_id, f, t))
    return _f(op.get("dr")) - _f(op.get("cr")), rows


# ════════════════════════════════════════════════════════════
#  REGISTRY
# ════════════════════════════════════════════════════════════
REGISTRY = OrderedDict()


def report(rid, group, title, params=("period",), desc="", needs=None):
    """needs (v7.0): who sees the report — a business module key
    ('restaurant', 'production' …), a country 'country:IN' / 'country:AE'
    / 'country:SA', or None for everyone."""
    def deco(fn):
        REGISTRY[rid] = {"id": rid, "group": group, "title": title,
                         "params": tuple(params), "desc": desc, "fn": fn, "needs": needs}
        return fn
    return deco


def shown(spec):
    """Does this company use the report? (its country and business modules)"""
    need = spec.get("needs")
    if not need:
        return True
    try:
        if str(need).startswith("country:"):
            from modules import country as _ctry
            return _ctry.code() in str(need)[8:].upper().split(",")
        from modules import business as _biz
        return any(_biz.has(k.strip()) for k in str(need).split(","))
    except Exception:
        return True


def catalog(all_reports=False):
    return [{k: v for k, v in r.items() if k != "fn"} for r in REGISTRY.values()
            if all_reports or shown(r)]


def run(rid, **kw):
    if rid not in REGISTRY:
        raise KeyError(f"Unknown report '{rid}'")
    spec = REGISTRY[rid]
    today = _dt.date.today()
    f = kw.get("date_from") or today.replace(day=1).isoformat()
    t = kw.get("date_to") or today.isoformat()
    if f > t:
        f, t = t, f
    kw["date_from"], kw["date_to"] = f, t
    res = spec["fn"](**kw)
    try:
        product_columns(res, rid)
    except Exception:
        pass
    res.setdefault("id", rid)
    res.setdefault("group", spec["group"])
    if not res.get("subtitle"):
        res["subtitle"] = (f"As of {t}" if "as_of" in spec["params"]
                           else _period_text(f, t))
    _localize(res)
    return res


def _localize(res):
    """v7.0: report words in the company's country (SAR → AED / ₹, VAT →
    GST) — on screen, in Excel / PDF exports and for the connector."""
    try:
        from modules import country as _K
        if _K.is_sa():
            return
        L = _K.localize
        for k in ("title", "subtitle"):
            if isinstance(res.get(k), str):
                res[k] = L(res[k])
        for c_ in res.get("columns") or []:
            if isinstance(c_, dict) and isinstance(c_.get("label"), str):
                c_["label"] = L(c_["label"])
        res["notes"] = [L(n) if isinstance(n, str) else n for n in (res.get("notes") or [])]
        res["summary"] = [((L(s[0]),) + tuple(s[1:])) if isinstance(s, (list, tuple)) and s
                          and isinstance(s[0], str) else s for s in (res.get("summary") or [])]
    except Exception:
        pass


# ── v6.9.4: product reports offer more columns ────────────────
# (columns tool ▸ ➕ More columns: Profit, Profit %, Cost, PLU, Stock,
# Unit 1-3, Multi barcode …). rid: (qty key, net key, unit key, mode,
# extra columns the report already has).
PRODUCT_REPORTS = {
    "sales_items": ("qty", "net_sales", None, "sale",
                    {"x_profit", "x_profit_pct", "x_cost", "x_category"}),
    "sales_lines": ("qty", "net", "unit", "sale", {"x_name_ar"}),
    "purchase_items": ("qty", "net", None, "purchase", set()),
    "purchase_lines": ("qty", "net", "unit", "purchase", set()),
    "stock": ("qty", None, None, "stock",
              {"x_cost", "x_price", "x_stock", "x_stock_value", "x_reorder", "x_category"}),
    "stock_movement": ("closing", None, None, "stock", set()),
    "low_stock": ("qty", None, None, "stock", {"x_cost", "x_stock", "x_category"}),
    "dead_stock": ("qty", None, None, "stock", {"x_stock", "x_category"}),
    "stock_by_branch": ("total", None, None, "stock", set()),
    "expiry_report": ("qty", None, None, "stock", set()),
    "stock_transactions": (None, None, None, "unit", set()),
}
_KIND = {"money": "money", "pct": "pct", "qty": "qty", "text": "text"}


def product_columns(res, rid):
    """Add the extra product columns (hidden until the user ticks them)
    to a product report: res["extra_columns"] lists them."""
    spec = PRODUCT_REPORTS.get(rid)
    rows = res.get("rows") or []
    if not spec or not rows:
        return res
    from modules import product_info as PI
    qk, nk, uk, mode, skip = spec
    flds = [f for f in PI.fields() if f[0] not in skip]
    have = {c["key"] for c in res["columns"]}
    flds = [f for f in flds if f[0] not in have]
    if not flds:
        return res
    tp = ts = tv = 0.0
    hit = False
    for r in rows:
        if r.get("_style") or not r.get("code"):
            continue
        v = PI.values(r.get("code"), qty=r.get(qk) if qk else None,
                      net=r.get(nk) if nk else None, unit=r.get(uk) if uk else None,
                      mode="purchase" if mode == "purchase" else "sale")
        if not v:
            continue
        hit = True
        for k, *_x in flds:
            r[k] = v.get(k)
        tp += v.get("x_profit") or 0.0
        ts += v.get("_sale") or 0.0
        tv += v.get("x_stock_value") or 0.0
    if not hit:
        return res
    res["columns"] = list(res["columns"]) + [col(k, lbl, _KIND.get(kind, "text"), w)
                                             for k, lbl, kind, w, _s in flds]
    res["extra_columns"] = [f[0] for f in flds]
    t = res.get("totals")
    if t is not None and qk:                       # per-unit figures have no total
        t["x_profit"] = round(tp, 2)
        t["x_profit_pct"] = round(tp / ts * 100.0, 2) if ts else None
        t["x_stock_value"] = round(tv, 2)
    return res


GROUP_ACC = "Accounts & VAT"
GROUP_SALES = "Sales & POS"
GROUP_PUR = "Purchases"
GROUP_INV = "Inventory"


# ════════════════════════════════════════════════════════════
#  ACCOUNTS & VAT
# ════════════════════════════════════════════════════════════
def _vat_cat(code, vcode, rate):
    c = str(vcode or code or "").strip().upper()
    if c in ("S", "Z", "E", "O"):
        return c
    return "S" if _f(rate) > 0 else "Z"


def _foreign(code):
    c = str(code or "").strip().upper()
    return bool(c) and c not in ("SA", "KSA", "SAU")


EXPORT_VATEX = ("VATEX-SA-32", "VATEX-SA-33")
CITIZEN_VATEX = ("VATEX-SA-HEA", "VATEX-SA-EDU", "VATEX-SA-MLTRY")


def previous_credit(date_from):
    """VAT credit to carry into this period (box 15): the last saved return
    before it that ended in a refund position and was not refunded."""
    try:
        r = _q1("SELECT TOP 1 net_vat_payable, return_status FROM vat_returns "
                "WHERE period_to < ? ORDER BY period_to DESC", (date_from,)) or {}
    except Exception:
        return 0.0
    n = _f(r.get("net_vat_payable"))
    if n < 0 and str(r.get("return_status") or "").lower() not in ("refunded", "paid"):
        return round(-n, 2)
    return 0.0


def vat_return_data(date_from, date_to, corrections=0.0, credit_bf=None,
                    reverse_charge=0.0, rc_not_recoverable=0.0):
    """ZATCA VAT return boxes for a period, INCLUDING sales returns,
    purchase returns and expense VAT.
    v6.0: box 2 (citizen health / education), box 4 (exports to customers
    outside KSA), box 8 (purchases from foreign suppliers with VAT paid at
    customs), box 9 (reverse charge — typed in), box 14 (corrections typed
    in), box 15 (credit carried forward — last refund position by default).
    Out-of-scope sales are NOT part of the return (they were in box 5)."""
    f, t = date_from, date_to
    cats = {c: {"amt": 0.0, "adj": 0.0, "vat": 0.0, "adj_vat": 0.0} for c in "SZEO"}
    special = {k: {"amt": 0.0, "adj": 0.0} for k in ("citizen", "export")}
    has_ex = True
    try:
        rows = _q("SELECT sil.vat_category_code AS code, vc.category_code AS vcode, "
                  "sil.vat_rate, sil.line_amount, sil.vat_amount, si.invoice_type_code, "
                  "si.customer_country_code AS ctry, sil.vat_exemption_code AS vatex "
                  "FROM sales_invoice_lines sil JOIN sales_invoices si ON si.invoice_id=sil.invoice_id "
                  "LEFT JOIN vat_categories vc ON vc.vat_category_id=sil.vat_category_id "
                  "WHERE si.invoice_status='confirmed' AND si.invoice_date BETWEEN ? AND ?", (f, t))
    except Exception:
        has_ex = False
        rows = _q("SELECT sil.vat_category_code AS code, vc.category_code AS vcode, "
                  "sil.vat_rate, sil.line_amount, sil.vat_amount, si.invoice_type_code "
                  "FROM sales_invoice_lines sil JOIN sales_invoices si ON si.invoice_id=sil.invoice_id "
                  "LEFT JOIN vat_categories vc ON vc.vat_category_id=sil.vat_category_id "
                  "WHERE si.invoice_status='confirmed' AND si.invoice_date BETWEEN ? AND ?", (f, t))
    for r in rows:
        c = _vat_cat(r.get("code"), r.get("vcode"), r.get("vat_rate"))
        credit = str(r.get("invoice_type_code") or "388") == "381"      # credit note
        vx = str(r.get("vatex") or "").upper() if has_ex else ""
        if c == "Z" and (vx in CITIZEN_VATEX):
            special["citizen"]["adj" if credit else "amt"] += -_f(r["line_amount"]) if credit \
                else _f(r["line_amount"])
            continue
        if c == "Z" and (vx in EXPORT_VATEX or (has_ex and _foreign(r.get("ctry")))):
            special["export"]["adj" if credit else "amt"] += -_f(r["line_amount"]) if credit \
                else _f(r["line_amount"])
            continue
        if credit:
            cats[c]["adj"] -= _f(r["line_amount"]); cats[c]["adj_vat"] -= _f(r["vat_amount"])
        else:
            cats[c]["amt"] += _f(r["line_amount"]); cats[c]["vat"] += _f(r["vat_amount"])
    sret = {"n": 0, "amt": 0.0, "vat": 0.0}
    for r in _return_lines("sales", f, t):
        c = _vat_cat(r.get("vat_category_code"), None, r.get("vat_rate"))
        cats[c]["adj"] -= _f(r["line_amount"]); cats[c]["adj_vat"] -= _f(r["vat_amount"])
        sret["amt"] += _f(r["line_amount"]); sret["vat"] += _f(r["vat_amount"])
    pur = {c: {"amt": 0.0, "adj": 0.0, "vat": 0.0, "adj_vat": 0.0} for c in "SZE"}
    imp = {"amt": 0.0, "vat": 0.0}
    try:
        prow = _q("SELECT pil.vat_rate, pil.line_amount, pil.vat_amount, vc.category_code AS vcode, "
                  "s.country_code AS ctry FROM purchase_invoice_lines pil JOIN purchase_invoices pi "
                  "ON pi.purchase_id=pil.purchase_id LEFT JOIN suppliers s ON s.supplier_id=pi.supplier_id "
                  "LEFT JOIN vat_categories vc ON vc.vat_category_id=pil.vat_category_id "
                  "WHERE pi.invoice_status='confirmed' AND pi.invoice_date BETWEEN ? AND ?", (f, t))
    except Exception:
        prow = _q("SELECT pil.vat_rate, pil.line_amount, pil.vat_amount, vc.category_code AS vcode "
                  "FROM purchase_invoice_lines pil JOIN purchase_invoices pi ON pi.purchase_id=pil.purchase_id "
                  "LEFT JOIN vat_categories vc ON vc.vat_category_id=pil.vat_category_id "
                  "WHERE pi.invoice_status='confirmed' AND pi.invoice_date BETWEEN ? AND ?", (f, t))
    for r in prow:
        c = _vat_cat(None, r.get("vcode"), r.get("vat_rate"))
        c = c if c in pur else "E"
        if c == "S" and _foreign(r.get("ctry")):        # import: VAT paid at customs
            imp["amt"] += _f(r["line_amount"]); imp["vat"] += _f(r["vat_amount"])
            continue
        pur[c]["amt"] += _f(r["line_amount"]); pur[c]["vat"] += _f(r["vat_amount"])
    pret = {"amt": 0.0, "vat": 0.0}
    for r in _return_lines("purchase", f, t):
        c = _vat_cat(r.get("vat_category_code"), None, r.get("vat_rate"))
        c = c if c in pur else "E"
        pur[c]["adj"] -= _f(r["line_amount"]); pur[c]["adj_vat"] -= _f(r["vat_amount"])
        pret["amt"] += _f(r["line_amount"]); pret["vat"] += _f(r["vat_amount"])
    exp = {"amt": 0.0, "vat": 0.0}
    if _has_table("expenses"):
        e = _q1("SELECT ISNULL(SUM(amount),0) AS amt, ISNULL(SUM(vat_amount),0) AS vat "
                "FROM expenses WHERE ISNULL(vat_amount,0) <> 0 AND expense_date BETWEEN ? AND ?",
                (f, t)) or {}
        exp = {"amt": _f(e.get("amt")), "vat": _f(e.get("vat"))}
        pur["S"]["amt"] += exp["amt"]; pur["S"]["vat"] += exp["vat"]

    def R(x):
        return round(x, 2) + 0.0          # + 0.0: never show "-0.00"

    boxes = []

    def box(no, label, a, adj, v, style=""):
        boxes.append({"box": no, "label": label, "amount": R(a), "adjustment": R(adj),
                      "vat": R(v), "_style": style})

    S, Z, E, O = cats["S"], cats["Z"], cats["E"], cats["O"]
    CZ, EX = special["citizen"], special["export"]
    box("1", "Standard rated sales (15%)", S["amt"], S["adj"], S["vat"] + S["adj_vat"], "hi")
    box("2", "Special sales to citizens (health / education)", CZ["amt"], CZ["adj"], 0)
    box("3", "Zero rated domestic sales", Z["amt"], Z["adj"], 0)
    box("4", "Exports", EX["amt"], EX["adj"], 0)
    box("5", "Exempt sales", E["amt"], E["adj"], 0)
    out_amt = S["amt"] + Z["amt"] + E["amt"] + CZ["amt"] + EX["amt"]
    out_adj = S["adj"] + Z["adj"] + E["adj"] + CZ["adj"] + EX["adj"]
    output_vat = S["vat"] + S["adj_vat"] + Z["vat"] + Z["adj_vat"]
    box("6", "Total sales", out_amt, out_adj, output_vat, "total")
    PS, PZ, PE = pur["S"], pur["Z"], pur["E"]
    box("7", "Standard rated domestic purchases (incl. expenses)", PS["amt"], PS["adj"],
        PS["vat"] + PS["adj_vat"], "hi")
    box("8", "Imports subject to VAT paid at customs", imp["amt"], 0, imp["vat"])
    rc_amt, rc_nr = _f(reverse_charge), _f(rc_not_recoverable)
    rc_vat = round(rc_nr * 0.15, 2)            # VAT on the part that can't be recovered
    box("9", "Imports subject to VAT (reverse charge)", rc_amt, rc_nr, rc_vat)
    box("10", "Zero rated purchases", PZ["amt"], PZ["adj"], 0)
    box("11", "Exempt purchases", PE["amt"], PE["adj"], 0)
    in_amt = PS["amt"] + PZ["amt"] + PE["amt"] + imp["amt"] + rc_amt
    in_adj = PS["adj"] + PZ["adj"] + PE["adj"]
    input_vat = PS["vat"] + PS["adj_vat"] + imp["vat"]
    box("12", "Total purchases", in_amt, in_adj, input_vat, "total")
    net = output_vat - input_vat + rc_vat
    box("13", "Total VAT due for current period", 0, 0, net, "total")
    corr = round(_f(corrections), 2)
    cbf = previous_credit(f) if credit_bf is None else round(_f(credit_bf), 2)
    box("14", "Corrections from previous period (±5,000)", 0, 0, corr)
    box("15", "VAT credit carried forward", 0, 0, -cbf)
    due = round(net + corr - cbf, 2)
    box("16", "Net VAT due (payable / refundable)", 0, 0, due, "grand")
    return {
        "boxes": boxes, "output_vat": R(output_vat), "input_vat": R(input_vat),
        "net": R(due), "net_period": R(net), "corrections": corr, "credit_bf": cbf,
        "out_of_scope": R(O["amt"] + O["adj"]), "imports": {k: R(v) for k, v in imp.items()},
        "sales_amount": R(out_amt + out_adj),
        "purchase_amount": R(in_amt + in_adj),
        "sales_returns": {k: R(v) for k, v in sret.items()},
        "purchase_returns": {k: R(v) for k, v in pret.items()},
        "expenses": {k: R(v) for k, v in exp.items()},
        "credit_notes_vat": R(-sum(c["adj_vat"] for c in cats.values()) - sret["vat"]),
        # legacy keys used by the saved vat_returns table
        "standard": {"taxable": R(S["amt"] + S["adj"]), "vat": R(S["vat"] + S["adj_vat"])},
        "zero": {"taxable": R(Z["amt"] + Z["adj"]), "vat": 0.0},
        "exempt": {"taxable": R(E["amt"] + E["adj"] + O["amt"] + O["adj"]), "vat": 0.0},
        "purchases": {"taxable": R(in_amt + in_adj), "vat": R(input_vat)},
    }


@report("vat_return", GROUP_ACC, "VAT Return",
        desc="Output VAT on sales less returns, input VAT on purchases and "
             "expenses less purchase returns, net VAT payable.",
        needs="country:SA")
def r_vat_return(date_from, date_to, **_):
    d = vat_return_data(date_from, date_to)
    cols = [col("box", "Box", width=50), col("label", "Description", width=330),
            col("amount", "Amount (SAR)", "money"), col("adjustment", "Adjustment", "money"),
            col("vat", "VAT (SAR)", "money")]
    summ = [("Output VAT", d["output_vat"], "money"), ("Input VAT", d["input_vat"], "money"),
            ("Net VAT " + ("payable" if d["net"] >= 0 else "refundable"), abs(d["net"]), "money")]
    notes = [f"Sales returns deducted: {d['sales_returns']['amt']:,.2f} "
             f"(VAT {d['sales_returns']['vat']:,.2f}).",
             f"Purchase returns deducted: {d['purchase_returns']['amt']:,.2f} "
             f"(VAT {d['purchase_returns']['vat']:,.2f}).",
             f"Expenses with VAT included in box 7: {d['expenses']['amt']:,.2f} "
             f"(VAT {d['expenses']['vat']:,.2f}).",
             f"Out-of-scope sales (not in the return): {d.get('out_of_scope', 0):,.2f}.",
             "Box 9 / 14 / 15: type them in the VAT Return form (Reports ▸ VAT Report). "
             "Box 15 here = the last saved return's refund position.",
             "Verify with source documents before filing on the ZATCA portal."]
    rows = [dict(b) for b in d["boxes"]]
    for b in rows:
        if b["box"] in ("13", "14", "15", "16"):
            b["amount"] = b["adjustment"] = None
    return _result("VAT Return", cols, rows, None, summ, notes)


@report("vat_detail", GROUP_ACC, "VAT Detail Register",
        desc="Every sales, return, purchase and expense document with its VAT.")
def r_vat_detail(date_from, date_to, **_):
    f, t = date_from, date_to
    rows = []
    for h in _sales_headers(f, t):
        cn = _is_credit_note(h)
        sg = -1 if cn else 1
        rows.append({"date": _d(h["invoice_date"]), "type": "Credit note" if cn else "Sale",
                     "number": h["invoice_number"], "party": h.get("customer_name_en"),
                     "vat_no": h.get("customer_vat_no") or "",
                     "taxable": sg * _f(h.get("taxable_amount")),
                     "out_vat": sg * _f(h.get("vat_amount")), "in_vat": 0.0})
    for h in _return_headers("sales", f, t):
        rows.append({"date": _d(h["return_date"]), "type": "Sales return",
                     "number": h["return_number"], "party": h.get("customer_name_en"),
                     "vat_no": h.get("customer_vat_no") or "",
                     "taxable": -_f(h.get("taxable_amount")),
                     "out_vat": -_f(h.get("vat_amount")), "in_vat": 0.0})
    for h in _purchase_headers(f, t):
        rows.append({"date": _d(h["invoice_date"]), "type": "Purchase",
                     "number": h["invoice_number"] + (f" / {h['supplier_invoice_no']}"
                                                      if h.get("supplier_invoice_no") else ""),
                     "party": h.get("supplier_name_en"), "vat_no": h.get("supplier_vat_no") or "",
                     "taxable": _f(h.get("taxable_amount")), "out_vat": 0.0,
                     "in_vat": _f(h.get("vat_amount"))})
    for h in _return_headers("purchase", f, t):
        rows.append({"date": _d(h["return_date"]), "type": "Purchase return",
                     "number": h["return_number"], "party": h.get("supplier_name_en"),
                     "vat_no": h.get("supplier_vat_no") or "",
                     "taxable": -_f(h.get("taxable_amount")), "out_vat": 0.0,
                     "in_vat": -_f(h.get("vat_amount"))})
    if _has_table("expenses"):
        for e in _q("SELECT e.*, a.account_name_en FROM expenses e LEFT JOIN accounts a "
                    "ON a.account_id=e.expense_account_id WHERE ISNULL(e.vat_amount,0) <> 0 "
                    "AND e.expense_date BETWEEN ? AND ?", (f, t)):
            rows.append({"date": _d(e["expense_date"]), "type": "Expense",
                         "number": e.get("reference") or f"EXP-{e['expense_id']}",
                         "party": e.get("account_name_en") or e.get("description") or "",
                         "vat_no": "", "taxable": _f(e.get("amount")), "out_vat": 0.0,
                         "in_vat": _f(e.get("vat_amount"))})
    rows.sort(key=lambda r: (r["date"], r["type"]))
    cols = [col("date", "Date", width=90), col("type", "Type", width=110),
            col("number", "Document #", width=130), col("party", "Party", width=200),
            col("vat_no", "VAT No.", width=140), col("taxable", "Taxable", "money"),
            col("out_vat", "Output VAT", "money"), col("in_vat", "Input VAT", "money")]
    tot = _sum_totals(rows, ("taxable", "out_vat", "in_vat"), "party", "TOTAL")
    net = tot["out_vat"] - tot["in_vat"]
    return _result("VAT Detail Register", cols, rows, tot,
                   [("Output VAT", tot["out_vat"], "money"),
                    ("Input VAT", tot["in_vat"], "money"),
                    ("Net VAT", net, "money")])


# ── Output VAT / Input VAT registers (v6.2) ────────────────────
def _doc_kind(h):
    tc = str(h.get("invoice_type_code") or "388")
    if tc == "381":
        return "Credit note"
    if tc == "383":
        return "Debit note"
    if str(h.get("sale_channel") or "").upper() == "POS":
        return "POS (simplified)"
    return "Tax invoice" if str(h.get("customer_vat_no") or "").strip() else "Simplified"


def _vat_split():
    return {"std": 0.0, "zero": 0.0, "exempt": 0.0, "oos": 0.0, "vat": 0.0, "odd_vat": 0.0}


def _add_line(a, cat, amount, vat, vat_cats=("S", "Z")):
    """One document line into its VAT column. VAT counts like the VAT
    Return (sales: standard + zero rated lines, purchases: standard) —
    VAT on other lines is a data error, kept apart in odd_vat."""
    key = {"S": "std", "Z": "zero", "E": "exempt", "O": "oos"}.get(cat, "exempt")
    a[key] += amount
    if cat in vat_cats:
        a["vat"] += vat
    elif abs(vat) > 0.004:
        a["odd_vat"] += vat


def _sales_vat_lines(f, t):
    by = {}
    for r in _q("SELECT sil.invoice_id AS doc, sil.vat_category_code AS code, "
                "vc.category_code AS vcode, sil.vat_rate, sil.line_amount, sil.vat_amount "
                "FROM sales_invoice_lines sil JOIN sales_invoices si ON si.invoice_id=sil.invoice_id "
                "LEFT JOIN vat_categories vc ON vc.vat_category_id=sil.vat_category_id "
                "WHERE si.invoice_status='confirmed' AND si.invoice_date BETWEEN ? AND ?", (f, t)):
        _add_line(by.setdefault(r["doc"], _vat_split()),
                  _vat_cat(r.get("code"), r.get("vcode"), r.get("vat_rate")),
                  _f(r["line_amount"]), _f(r["vat_amount"]))
    return by


def _purchase_vat_lines(f, t):
    by = {}
    for r in _q("SELECT pil.purchase_id AS doc, vc.category_code AS vcode, pil.vat_rate, "
                "pil.line_amount, pil.vat_amount FROM purchase_invoice_lines pil "
                "JOIN purchase_invoices pi ON pi.purchase_id=pil.purchase_id "
                "LEFT JOIN vat_categories vc ON vc.vat_category_id=pil.vat_category_id "
                "WHERE pi.invoice_status='confirmed' AND pi.invoice_date BETWEEN ? AND ?", (f, t)):
        c = _vat_cat(None, r.get("vcode"), r.get("vat_rate"))
        _add_line(by.setdefault(r["doc"], _vat_split()), c if c in "SZE" else "E",
                  _f(r["line_amount"]), _f(r["vat_amount"]), ("S",))
    return by


def _return_vat_lines(kind, f, t):
    by = {}
    for r in _return_lines(kind, f, t):
        c = _vat_cat(r.get("vat_category_code"), None, r.get("vat_rate"))
        if kind == "purchase" and c not in "SZE":
            c = "E"
        _add_line(by.setdefault(r["return_id"], _vat_split()), c,
                  _f(r["line_amount"]), _f(r["vat_amount"]),
                  ("S",) if kind == "purchase" else ("S", "Z"))
    return by


def _vat_row(base, a, sign=1):
    row = dict(base)
    for k in ("std", "zero", "exempt", "oos", "vat"):
        row[k] = round(sign * a.get(k, 0.0), 2) + 0.0
    row["taxable"] = round(row["std"] + row["zero"] + row["exempt"] + row["oos"], 2) + 0.0
    row["total"] = round(row["taxable"] + row["vat"], 2) + 0.0
    return row


@report("output_vat", GROUP_ACC, "Output VAT Register (Sales VAT)", params=("period", "party"),
        desc="Every sales invoice, POS bill, credit / debit note and sales return with "
             "its standard (15%), zero-rated, exempt amounts and output VAT. "
             "Total output VAT = VAT Return box 6.")
def r_output_vat(date_from, date_to, party="", **_):
    f, t = date_from, date_to
    rows, odd = [], 0.0
    lines = _sales_vat_lines(f, t)
    for h in _sales_headers(f, t):
        if not _match(party, h.get("customer_name_en"), h.get("invoice_number"),
                      h.get("customer_vat_no")):
            continue
        a = lines.get(h["invoice_id"]) or _vat_split()
        sg = -1 if _is_credit_note(h) else 1
        odd += a["odd_vat"]
        rows.append(_vat_row({"date": _d(h["invoice_date"]), "type": _doc_kind(h),
                              "number": h["invoice_number"],
                              "ref": h.get("original_invoice_no") or "",
                              "party": h.get("customer_name_en") or "",
                              "vat_no": h.get("customer_vat_no") or ""}, a, sg))
    rl = _return_vat_lines("sales", f, t)
    for h in _return_headers("sales", f, t):
        if not _match(party, h.get("customer_name_en"), h.get("return_number"),
                      h.get("customer_vat_no")):
            continue
        a = rl.get(h["return_id"]) or _vat_split()
        odd += a["odd_vat"]
        rows.append(_vat_row({"date": _d(h["return_date"]), "type": "Sales return",
                              "number": h["return_number"],
                              "ref": h.get("original_invoice_no") or "",
                              "party": h.get("customer_name_en") or "",
                              "vat_no": h.get("customer_vat_no") or ""}, a, -1))
    rows.sort(key=lambda r: (r["date"], str(r["number"])))
    cols = [col("date", "Date", width=90), col("type", "Type", width=120),
            col("number", "Document #", width=135), col("ref", "Orig. invoice", width=110),
            col("party", "Customer", width=190), col("vat_no", "Customer VAT No.", width=140),
            col("std", "Standard 15%", "money"), col("zero", "Zero rated", "money"),
            col("exempt", "Exempt", "money"), col("oos", "Out of scope", "money"),
            col("taxable", "Taxable total", "money"), col("vat", "Output VAT", "money"),
            col("total", "Total incl. VAT", "money")]
    tot = _sum_totals(rows, ("std", "zero", "exempt", "oos", "taxable", "vat", "total"),
                      "party", "TOTAL")
    notes = ["Credit notes and sales returns are minus. Output VAT here = VAT Return "
             "box 6 (VAT) when no customer / document filter is used.",
             "Out-of-scope sales are not part of the VAT return."]
    if abs(odd) > 0.004:
        notes.append(f"⚠ {odd:,.2f} VAT is on exempt / out-of-scope lines — "
                     "not counted (check the VAT category of those items).")
    return _result("Output VAT Register", cols, rows, tot,
                   [("Documents", len(rows), "int"), ("Taxable sales", tot["taxable"], "money"),
                    ("Standard 15%", tot["std"], "money"), ("Zero rated", tot["zero"], "money"),
                    ("Exempt", tot["exempt"], "money"), ("Output VAT", tot["vat"], "money")],
                   notes)


@report("input_vat", GROUP_ACC, "Input VAT Register (Purchase VAT)", params=("period", "party"),
        desc="Every purchase invoice, purchase return and expense with VAT: supplier "
             "VAT No., standard / zero-rated / exempt amounts and input VAT. "
             "Total input VAT = VAT Return boxes 7 + 8.")
def r_input_vat(date_from, date_to, party="", **_):
    f, t = date_from, date_to
    rows, odd = [], 0.0
    lines = _purchase_vat_lines(f, t)
    try:
        ctry = {r["supplier_id"]: r.get("country_code") for r in
                _q("SELECT supplier_id, country_code FROM suppliers")}
    except Exception:
        ctry = {}
    for h in _purchase_headers(f, t):
        if not _match(party, h.get("supplier_name_en"), h.get("invoice_number"),
                      h.get("supplier_invoice_no"), h.get("supplier_vat_no")):
            continue
        a = lines.get(h["purchase_id"]) or _vat_split()
        odd += a["odd_vat"]
        imp = _foreign(ctry.get(h.get("supplier_id")))
        rows.append(_vat_row({"date": _d(h["invoice_date"]),
                              "type": "Import" if imp else "Purchase",
                              "number": h["invoice_number"],
                              "sup_inv": h.get("supplier_invoice_no") or "",
                              "party": h.get("supplier_name_en") or "",
                              "vat_no": h.get("supplier_vat_no") or ""}, a))
    rl = _return_vat_lines("purchase", f, t)
    for h in _return_headers("purchase", f, t):
        if not _match(party, h.get("supplier_name_en"), h.get("return_number"),
                      h.get("supplier_vat_no")):
            continue
        a = rl.get(h["return_id"]) or _vat_split()
        odd += a["odd_vat"]
        rows.append(_vat_row({"date": _d(h["return_date"]), "type": "Purchase return",
                              "number": h["return_number"],
                              "sup_inv": h.get("original_purchase_no") or "",
                              "party": h.get("supplier_name_en") or "",
                              "vat_no": h.get("supplier_vat_no") or ""}, a, -1))
    if _has_table("expenses"):
        has_vno = _has_col("expenses", "supplier_vat_no")
        for e in _q("SELECT e.*, a.account_name_en FROM expenses e LEFT JOIN accounts a "
                    "ON a.account_id=e.expense_account_id WHERE ISNULL(e.vat_amount,0) <> 0 "
                    "AND e.expense_date BETWEEN ? AND ?", (f, t)):
            name = e.get("account_name_en") or e.get("description") or "Expense"
            if not _match(party, name, e.get("reference"), e.get("description")):
                continue
            a = _vat_split()
            _add_line(a, "S", _f(e.get("amount")), _f(e.get("vat_amount")))
            rows.append(_vat_row({"date": _d(e["expense_date"]), "type": "Expense",
                                  "number": e.get("reference") or f"EXP-{e['expense_id']}",
                                  "sup_inv": "", "party": name,
                                  "vat_no": (e.get("supplier_vat_no") or "") if has_vno else ""}, a))
    rows.sort(key=lambda r: (r["date"], str(r["number"])))
    cols = [col("date", "Date", width=90), col("type", "Type", width=110),
            col("number", "Document #", width=135), col("sup_inv", "Supplier inv. #", width=110),
            col("party", "Supplier", width=190), col("vat_no", "Supplier VAT No.", width=140),
            col("std", "Standard 15%", "money"), col("zero", "Zero rated", "money"),
            col("exempt", "Exempt", "money"), col("taxable", "Taxable total", "money"),
            col("vat", "Input VAT", "money"), col("total", "Total incl. VAT", "money")]
    tot = _sum_totals(rows, ("std", "zero", "exempt", "taxable", "vat", "total"),
                      "party", "TOTAL")
    imp_vat = round(sum(r["vat"] for r in rows if r["type"] == "Import"), 2)
    notes = ["Purchase returns are minus. Expenses with VAT are standard rated. "
             "Input VAT here = VAT Return boxes 7 + 8 (VAT) when no filter is used.",
             "Import = supplier outside Saudi Arabia (VAT paid at customs, box 8)."]
    if abs(odd) > 0.004:
        notes.append(f"⚠ {odd:,.2f} VAT is on zero-rated / exempt lines — not counted "
                     "(check the VAT category of those items).")
    return _result("Input VAT Register", cols, rows, tot,
                   [("Documents", len(rows), "int"), ("Purchases (taxable)", tot["taxable"], "money"),
                    ("Standard 15%", tot["std"], "money"), ("Imports VAT", imp_vat, "money"),
                    ("Input VAT", tot["vat"], "money")], notes)


@report("trial_balance", GROUP_ACC, "Trial Balance",
        desc="Opening, period movement and closing balance of every account.")
def r_trial_balance(date_from, date_to, **_):
    f, t = date_from, date_to
    data = _q("SELECT a.account_code, a.account_name_en, a.account_type, "
              "SUM(CASE WHEN je.entry_date < ? THEN jl.debit ELSE 0 END) AS odr, "
              "SUM(CASE WHEN je.entry_date < ? THEN jl.credit ELSE 0 END) AS ocr, "
              "SUM(CASE WHEN je.entry_date >= ? THEN jl.debit ELSE 0 END) AS pdr, "
              "SUM(CASE WHEN je.entry_date >= ? THEN jl.credit ELSE 0 END) AS pcr "
              "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
              "JOIN accounts a ON a.account_id=jl.account_id "
              "WHERE je.is_posted=1 AND je.entry_date <= ? "
              "GROUP BY a.account_code, a.account_name_en, a.account_type "
              "ORDER BY a.account_code", (f, f, f, f, t))
    rows = []
    for r in data:
        o = _f(r["odr"]) - _f(r["ocr"])
        c = o + _f(r["pdr"]) - _f(r["pcr"])
        if not any(abs(x) > 0.004 for x in (o, _f(r["pdr"]), _f(r["pcr"]), c)):
            continue
        rows.append({"code": r["account_code"], "name": r["account_name_en"],
                     "type": (r.get("account_type") or "").title(),
                     "op_dr": round(o, 2) if o > 0 else 0.0,
                     "op_cr": round(-o, 2) if o < 0 else 0.0,
                     "dr": round(_f(r["pdr"]), 2), "cr": round(_f(r["pcr"]), 2),
                     "cl_dr": round(c, 2) if c > 0 else 0.0,
                     "cl_cr": round(-c, 2) if c < 0 else 0.0})
    cols = [col("code", "Code", width=80), col("name", "Account", width=220),
            col("type", "Type", width=80),
            col("op_dr", "Opening Dr", "money"), col("op_cr", "Opening Cr", "money"),
            col("dr", "Debit", "money"), col("cr", "Credit", "money"),
            col("cl_dr", "Closing Dr", "money"), col("cl_cr", "Closing Cr", "money")]
    tot = _sum_totals(rows, ("op_dr", "op_cr", "dr", "cr", "cl_dr", "cl_cr"), "name", "TOTAL")
    ok = abs(tot["cl_dr"] - tot["cl_cr"]) < 0.01
    return _result("Trial Balance", cols, rows, tot,
                   [("Closing debit", tot["cl_dr"], "money"),
                    ("Closing credit", tot["cl_cr"], "money"),
                    ("Balanced", "Yes ✔" if ok else "NO — difference %.2f"
                     % (tot["cl_dr"] - tot["cl_cr"]), "text")])


def _pl_rows(f, t):
    try:
        from modules import account_defaults as ad
        cos_ids = {ad.cogs(), ad.purchase_return(), ad.stock_adjust()}
    except Exception:
        cos_ids = set()
    data = _q("SELECT a.account_id, a.account_code, a.account_name_en, a.account_type, "
              "ISNULL(SUM(jl.debit),0) AS dr, ISNULL(SUM(jl.credit),0) AS cr "
              "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
              "JOIN accounts a ON a.account_id=jl.account_id "
              "WHERE je.is_posted=1 AND je.entry_date BETWEEN ? AND ? "
              "AND a.account_type IN ('revenue','expense') AND ISNULL(je.source_type,'')<>'year_close' "
              "GROUP BY a.account_id, a.account_code, a.account_name_en, a.account_type "
              "ORDER BY a.account_code", (f, t))
    rev, cos, exp = [], [], []
    for r in data:
        if r["account_type"] == "revenue":
            rev.append((r["account_code"], r["account_name_en"], _f(r["cr"]) - _f(r["dr"])))
        elif r["account_id"] in cos_ids or str(r["account_code"]).startswith("51"):
            cos.append((r["account_code"], r["account_name_en"], _f(r["dr"]) - _f(r["cr"])))
        else:
            exp.append((r["account_code"], r["account_name_en"], _f(r["dr"]) - _f(r["cr"])))
    return rev, cos, exp


@report("profit_loss", GROUP_ACC, "Profit & Loss",
        desc="Net sales, cost of sales, gross profit, expenses and net profit.")
def r_profit_loss(date_from, date_to, **_):
    rev, cos, exp = _pl_rows(date_from, date_to)
    tr = sum(a for _, _, a in rev); tc = sum(a for _, _, a in cos)
    te = sum(a for _, _, a in exp)
    gp = tr - tc; net = gp - te
    rows = []

    def pct(a):
        return round(100.0 * a / tr, 1) if tr else 0.0

    def head(label):
        rows.append({"code": "", "name": label, "amount": None, "pct": None, "_style": "head"})

    def line(code, name, a):
        rows.append({"code": code, "name": "    " + name, "amount": round(a, 2), "pct": pct(a)})

    def total(label, a, style="total"):
        rows.append({"code": "", "name": label, "amount": round(a, 2), "pct": pct(a), "_style": style})

    head("REVENUE")
    for c, n, a in rev:
        line(c, n, a)
    total("Net Revenue", tr)
    head("COST OF SALES")
    for c, n, a in cos:
        line(c, n, a)
    total("Total Cost of Sales", tc)
    total("GROSS PROFIT", gp, "grand")
    head("OPERATING EXPENSES")
    for c, n, a in exp:
        line(c, n, a)
    total("Total Expenses", te)
    total("NET PROFIT" if net >= 0 else "NET LOSS", net, "grand")
    cols = [col("code", "Code", width=80), col("name", "Particulars", width=320),
            col("amount", "Amount (SAR)", "money", 140), col("pct", "% of Revenue", "pct", 100)]
    notes = []
    if tr and not cos:
        notes.append("No cost of sales in this period — sales saved before v5.7 have no "
                     "cost entries. Use 'Post missing cost of sales' on the Reports screen.")
    return _result("Profit & Loss Statement", cols, rows, None,
                   [("Net revenue", tr, "money"), ("Gross profit", gp, "money"),
                    ("Net profit", net, "money")], notes)


@report("balance_sheet", GROUP_ACC, "Balance Sheet", params=("as_of",),
        desc="Statement of financial position: current & fixed assets, liabilities "
             "(VAT net), equity with profit to date — as of a date.")
def r_balance_sheet(date_to, **_):
    """Standard layout. Customer / supplier sub-ledgers are rolled into
    Accounts Receivable / Payable (see Customer / Supplier Balances for
    the list), VAT output and input are netted into one VAT line, and
    Fixed Assets are shown less accumulated depreciation."""
    data = _q("SELECT a.account_code, a.account_name_en, a.account_type, "
              "ISNULL(SUM(jl.debit),0) AS dr, ISNULL(SUM(jl.credit),0) AS cr "
              "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
              "JOIN accounts a ON a.account_id=jl.account_id "
              "WHERE je.is_posted=1 AND je.entry_date <= ? "
              "GROUP BY a.account_code, a.account_name_en, a.account_type "
              "ORDER BY a.account_code", (date_to,))
    names = {}
    try:
        for r in _q("SELECT account_code, account_name_en FROM accounts") or []:
            names[str(r["account_code"])] = r["account_name_en"]
    except Exception:
        pass

    cur_a, fix_a, cur_l, eq = {}, {}, {}, {}
    ar = ap = vat = profit = 0.0

    def add(bucket, code, name, amt):
        c, n, v = bucket.get(code, (code, name, 0.0))
        bucket[code] = (c, n, v + amt)

    for r in data:
        code = str(r["account_code"] or "")
        base = code.split("-")[0]
        dr, cr = _f(r["dr"]), _f(r["cr"])
        ty = r["account_type"]
        nm = r["account_name_en"] or code
        if ty == "revenue":
            profit += cr - dr
        elif ty == "expense":
            profit -= dr - cr
        elif base.startswith("12") and ty == "asset":
            ar += dr - cr
        elif base.startswith("21") and ty == "liability":
            ap += cr - dr
        elif base.startswith("22"):
            vat += cr - dr                      # output − input (+ = owed)
        elif ty == "asset" and base.startswith("15"):
            add(fix_a, base, names.get(base, nm) if base != code else nm, dr - cr)
        elif ty == "asset":
            add(cur_a, code, nm, dr - cr)
        elif ty == "liability":
            add(cur_l, code, nm, cr - dr)
        elif ty == "equity":
            add(eq, code, nm, cr - dr)
    if abs(ar) >= 0.005:
        cur_a["1200"] = ("1200", "Accounts receivable (customers)", ar)
    if vat < -0.005:
        cur_a["2200"] = ("2200", "VAT refundable (input more than output)", -vat)
    if abs(ap) >= 0.005:
        cur_l["2100"] = ("2100", "Accounts payable (suppliers)", ap)
    if vat > 0.005:
        cur_l["2200"] = ("2200", "VAT payable (output − input)", vat)

    rows = []

    def head(t, style="head"):
        rows.append({"code": "", "name": t, "amount": None, "_style": style})

    def items(bucket, indent="    "):
        tot = 0.0
        for c in sorted(bucket):
            code, n, a = bucket[c]
            if abs(a) < 0.005:
                continue
            rows.append({"code": code, "name": indent + n, "amount": round(a, 2)})
            tot += a
        return tot

    def total(t, v, style="total"):
        rows.append({"code": "", "name": t, "amount": round(v, 2), "_style": style})

    head("ASSETS")
    head("  Current assets")
    tca = items(cur_a, "      ")
    total("  Total current assets", tca)
    tfa = 0.0
    if any(abs(v[2]) >= 0.005 for v in fix_a.values()):
        head("  Fixed assets")
        tfa = items(fix_a, "      ")
        total("  Net fixed assets", tfa)
    ta = tca + tfa
    total("TOTAL ASSETS", ta, "grand")

    head("LIABILITIES")
    tl = items(cur_l)
    total("Total liabilities", tl)

    head("EQUITY")
    te = items(eq)
    rows.append({"code": "", "name": "    Profit / (loss) to date", "amount": round(profit, 2)})
    te += profit
    total("Total equity", te)
    total("TOTAL LIABILITIES + EQUITY", tl + te, "grand")
    diff = round(ta - (tl + te), 2)
    ok = abs(diff) < 0.01
    notes = []
    if not ok:
        notes.append(f"Difference {diff:,.2f} — an entry is not balanced. Check Trial Balance.")
    notes.append("Customer / supplier lists: Reports ▸ Customer Balances / Supplier Balances.")
    cols = [col("code", "Code", width=90), col("name", "Particulars", width=360),
            col("amount", "Amount (SAR)", "money", 150)]
    return _result("Balance Sheet", cols, rows, None,
                   [("Total assets", ta, "money"), ("Liabilities", tl, "money"),
                    ("Equity", te, "money"),
                    ("Balanced", "Yes ✔" if ok else f"NO ({diff:,.2f})", "text")], notes)


def _book(title, account, date_from, date_to, receipt_lbl="Receipt", pay_lbl="Payment"):
    acc = _account(account)
    cols = [col("date", "Date", width=90), col("jv", "Voucher", width=95),
            col("type", "Type", width=105), col("ref", "Ref", width=125),
            col("particulars", "Particulars", width=260),
            col("dr", receipt_lbl, "money"), col("cr", pay_lbl, "money"),
            col("balance", "Balance", "money")]
    if not acc:
        return _result(title, cols, [], None, [], ["Choose an account."])
    debit_normal = acc["account_type"] in ("asset", "expense")
    sgn = 1 if debit_normal else -1
    opening, lines = _acc_movement(acc["account_id"], date_from, date_to)
    bal = sgn * opening
    rows = [{"date": date_from, "jv": "", "type": "", "ref": "",
             "particulars": "Opening balance", "dr": None, "cr": None,
             "balance": round(bal, 2), "_style": "sub"}]
    tdr = tcr = 0.0
    for r in lines:
        dr, cr = _f(r["debit"]), _f(r["credit"])
        bal += sgn * (dr - cr)
        tdr += dr; tcr += cr
        part = r.get("description_en") or ""
        if r.get("line_desc") and r["line_desc"] not in part:
            part = f"{part} · {r['line_desc']}" if part else r["line_desc"]
        rows.append({"date": _d(r["entry_date"]), "jv": r["journal_number"],
                     "type": type_label(r["source_type"]), "ref": r.get("reference") or "",
                     "particulars": part, "dr": round(dr, 2) or None,
                     "cr": round(cr, 2) or None, "balance": round(bal, 2)})
    rows.append({"date": date_to, "jv": "", "type": "", "ref": "",
                 "particulars": "Closing balance", "dr": round(tdr, 2), "cr": round(tcr, 2),
                 "balance": round(bal, 2), "_style": "total"})
    return _result(title, cols, rows, None,
                   [("Account", f"{acc['account_code']} {acc['account_name_en']}", "text"),
                    ("Opening", sgn * opening, "money"), (receipt_lbl + "s", tdr, "money"),
                    (pay_lbl + "s", tcr, "money"), ("Closing", bal, "money")])


@report("cash_book", GROUP_ACC, "Cash Book", params=("period", "cashbank"),
        desc="Cash in / out with running balance.")
def r_cash_book(date_from, date_to, account=None, **_):
    return _book("Cash Book", account or _default_code("cash"), date_from, date_to)


@report("bank_book", GROUP_ACC, "Bank Book", params=("period", "cashbank"),
        desc="Bank / card deposits and payments with running balance.")
def r_bank_book(date_from, date_to, account=None, **_):
    return _book("Bank Book", account or _default_code("bank"), date_from, date_to,
                 "Deposit", "Withdrawal")


@report("ledger", GROUP_ACC, "General Ledger", params=("period", "account"),
        desc="Every posting to one account with running balance.")
def r_ledger(date_from, date_to, account=None, **_):
    return _book("General Ledger", account or _default_code("cash"), date_from, date_to,
                 "Debit", "Credit")


@report("day_book", GROUP_ACC, "Day Book",
        desc="All accounting entries of the period, voucher by voucher.")
def r_day_book(date_from, date_to, **_):
    data = _q("SELECT je.journal_id, je.entry_date, je.journal_number, je.source_type, "
              "je.reference, je.description_en, a.account_code, a.account_name_en, "
              "jl.description AS line_desc, jl.debit, jl.credit, %s AS voucher_type " % _vt_col() +
              "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
              "JOIN accounts a ON a.account_id=jl.account_id "
              "WHERE je.is_posted=1 AND je.entry_date BETWEEN ? AND ? "
              "ORDER BY je.entry_date, je.journal_id, jl.line_id", (date_from, date_to))
    rows, last = [], None
    for r in data:
        first = r["journal_id"] != last
        last = r["journal_id"]
        rows.append({"date": _d(r["entry_date"]) if first else "",
                     "jv": r["journal_number"] if first else "",
                     "type": vlabel(r) if first else "",
                     "ref": (r.get("reference") or "") if first else "",
                     "account": f"{r['account_code']}  {r['account_name_en']}",
                     "narration": r.get("line_desc") or (r.get("description_en") if first else ""),
                     "dr": round(_f(r["debit"]), 2) or None,
                     "cr": round(_f(r["credit"]), 2) or None})
    cols = [col("date", "Date", width=90), col("jv", "Voucher", width=105),
            col("type", "Vch Type", width=150), col("ref", "Ref", width=80),
            col("account", "Account", width=220), col("narration", "Narration", width=200),
            col("dr", "Debit", "money"), col("cr", "Credit", "money")]
    tot = _sum_totals(rows, ("dr", "cr"), "account", "TOTAL")
    n = len({r["journal_id"] for r in data})
    return _result("Day Book", cols, rows, tot,
                   [("Vouchers", n, "int"), ("Total debit", tot["dr"], "money"),
                    ("Total credit", tot["cr"], "money")])


@report("journal_register", GROUP_ACC, "Journal Register",
        desc="One line per voucher with amount and status.")
def r_journal_register(date_from, date_to, **_):
    data = _q("SELECT journal_number, entry_date, source_type, reference, description_en, "
              "total_debit, is_posted, %s AS voucher_type FROM journal_entries je "
              "WHERE entry_date BETWEEN ? AND ? "
              "ORDER BY entry_date, journal_id" % _vt_col(), (date_from, date_to))
    rows = [{"date": _d(r["entry_date"]), "jv": r["journal_number"],
             "type": vlabel(r), "ref": r.get("reference") or "",
             "narration": r.get("description_en") or "",
             "amount": round(_f(r["total_debit"]), 2),
             "status": "Posted" if r.get("is_posted") else "Draft"} for r in data]
    cols = [col("date", "Date", width=90), col("jv", "Voucher", width=105),
            col("type", "Vch Type", width=150), col("ref", "Ref", width=90),
            col("narration", "Narration", width=280), col("amount", "Amount", "money"),
            col("status", "Status", width=70)]
    return _result("Journal Register", cols, rows,
                   _sum_totals(rows, ("amount",), "narration", "TOTAL"),
                   [("Vouchers", len(rows), "int")])


@report("voucher_types", GROUP_ACC, "Voucher Type Summary",
        desc="How many SI, POS, SR, PI, PR, CR, BR, CP, BP, EX, JV … vouchers and their amount.")
def r_voucher_types(date_from, date_to, **_):
    from modules import voucher_types as _vt
    data = _q("SELECT source_type, total_debit, %s AS voucher_type FROM journal_entries je "
              "WHERE je.is_posted=1 AND entry_date BETWEEN ? AND ?" % _vt_col(),
              (date_from, date_to))
    by = OrderedDict((c, {"code": c, "name": n, "count": 0, "amount": 0.0}) for c, n in _vt.choices())
    for r in data:
        c = _vt.of_journal(r)
        d = by.setdefault(c, {"code": c, "name": _vt.name(c), "count": 0, "amount": 0.0})
        d["count"] += 1
        d["amount"] += _f(r["total_debit"])
    rows = [dict(d, amount=round(d["amount"], 2)) for d in by.values() if d["count"]]
    cols = [col("code", "Code", width=70), col("name", "Voucher type", width=220),
            col("count", "Vouchers", "int", 90), col("amount", "Amount", "money")]
    return _result("Voucher Type Summary", cols, rows,
                   _sum_totals(rows, ("count", "amount"), "name", "TOTAL"))


def _party_balances(kind, date_to):
    tbl, pid, code, name = (("customers", "customer_id", "customer_code", "customer_name_en")
                            if kind == "customer" else
                            ("suppliers", "supplier_id", "supplier_code", "supplier_name_en"))
    if not _has_col(tbl, "account_id"):
        return []
    try:
        from modules.aging import party_accounts
        ok = party_accounts(kind)
    except Exception:
        ok = None
    parties = [p for p in _q(f"SELECT * FROM {tbl} WHERE account_id IS NOT NULL")
               if ok is None or p["account_id"] in ok]
    bal = {r["account_id"]: _f(r["dr"]) - _f(r["cr"]) for r in _q(
        "SELECT jl.account_id, ISNULL(SUM(jl.debit),0) AS dr, ISNULL(SUM(jl.credit),0) AS cr "
        "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
        "WHERE je.is_posted=1 AND je.entry_date <= ? GROUP BY jl.account_id", (date_to,))}
    out = []
    for p in parties:
        b = bal.get(p["account_id"], 0.0)
        if kind == "supplier":
            b = -b
        if abs(b) < 0.005:
            continue
        out.append({"code": p.get(code) or "", "name": p.get(name) or "",
                    "phone": p.get("phone") or p.get("mobile") or "",
                    "vat_no": p.get("vat_registration_no") or "",
                    "due": round(b, 2) if b > 0 else 0.0,
                    "advance": round(-b, 2) if b < 0 else 0.0})
    out.sort(key=lambda r: -r["due"])
    # entries posted straight to the control account (no party chosen)
    try:
        from modules import account_defaults as ad
        ctrl = ad.receivable() if kind == "customer" else ad.payable()
        b = bal.get(ctrl, 0.0) * (1 if kind == "customer" else -1)
        if abs(b) >= 0.005:
            out.append({"code": "", "name": "(No party — posted to the control account)",
                        "phone": "", "vat_no": "",
                        "due": round(b, 2) if b > 0 else 0.0,
                        "advance": round(-b, 2) if b < 0 else 0.0})
    except Exception:
        pass
    return out


@report("receivables", GROUP_ACC, "Customer Balances", params=("as_of",),
        desc="What each customer owes (or has paid in advance).")
def r_receivables(date_to, **_):
    rows = _party_balances("customer", date_to)
    cols = [col("code", "Code", width=90), col("name", "Customer", width=240),
            col("phone", "Phone", width=110), col("vat_no", "VAT No.", width=140),
            col("due", "Owes you", "money"), col("advance", "Advance", "money")]
    tot = _sum_totals(rows, ("due", "advance"), "name", "TOTAL")
    return _result("Customer Balances", cols, rows, tot,
                   [("Customers", len(rows), "int"), ("Total receivable", tot["due"], "money"),
                    ("Advances", tot["advance"], "money")])


@report("payables", GROUP_ACC, "Supplier Balances", params=("as_of",),
        desc="What you owe each supplier.")
def r_payables(date_to, **_):
    rows = _party_balances("supplier", date_to)
    cols = [col("code", "Code", width=90), col("name", "Supplier", width=240),
            col("phone", "Phone", width=110), col("vat_no", "VAT No.", width=140),
            col("due", "You owe", "money"), col("advance", "Advance paid", "money")]
    tot = _sum_totals(rows, ("due", "advance"), "name", "TOTAL")
    return _result("Supplier Balances", cols, rows, tot,
                   [("Suppliers", len(rows), "int"), ("Total payable", tot["due"], "money")])


@report("expenses", GROUP_ACC, "Expense Report",
        desc="Expenses with VAT, grouped totals per expense account.")
def r_expenses(date_from, date_to, **_):
    if not _has_table("expenses"):
        return _result("Expense Report", [], [], None, [], ["Expenses table not installed."])
    data = _q("SELECT e.*, a.account_code, a.account_name_en FROM expenses e "
              "LEFT JOIN accounts a ON a.account_id=e.expense_account_id "
              "WHERE e.expense_date BETWEEN ? AND ? ORDER BY e.expense_date, e.expense_id",
              (date_from, date_to))
    rows, by = [], defaultdict(float)
    for e in data:
        amt, vat = _f(e.get("amount")), _f(e.get("vat_amount"))
        acc = f"{e.get('account_code') or ''} {e.get('account_name_en') or ''}".strip()
        by[acc] += amt
        rows.append({"date": _d(e["expense_date"]), "account": acc,
                     "desc": e.get("description") or "", "ref": e.get("reference") or "",
                     "method": (e.get("payment_method") or "").replace("_", " ").title(),
                     "amount": round(amt, 2), "vat": round(vat, 2), "total": round(amt + vat, 2)})
    cols = [col("date", "Date", width=90), col("account", "Expense account", width=200),
            col("desc", "Description", width=200), col("ref", "Ref", width=90),
            col("method", "Paid by", width=90), col("amount", "Amount", "money"),
            col("vat", "VAT", "money"), col("total", "Total", "money")]
    tot = _sum_totals(rows, ("amount", "vat", "total"), "desc", "TOTAL")
    summ = [(k, v, "money") for k, v in sorted(by.items(), key=lambda x: -x[1])[:6]]
    return _result("Expense Report", cols, rows, tot, summ)


@report("receipts_payments", GROUP_ACC, "Receipts & Payments",
        desc="Money received from customers and paid to suppliers.")
def r_receipts_payments(date_from, date_to, **_):
    data = _q("SELECT * FROM payments WHERE payment_date BETWEEN ? AND ? "
              "ORDER BY payment_date, payment_id", (date_from, date_to))
    rows = []
    for p in data:
        rec = str(p.get("payment_type") or "").lower() == "receipt"
        rows.append({"date": _d(p["payment_date"]), "number": p.get("payment_number") or "",
                     "type": "Receipt" if rec else "Payment",
                     "party": p.get("party_name_en") or "",
                     "method": (p.get("payment_method") or "").replace("_", " ").title(),
                     "ref": p.get("reference_number") or "",
                     "received": round(_f(p["amount"]), 2) if rec else None,
                     "paid": None if rec else round(_f(p["amount"]), 2)})
    cols = [col("date", "Date", width=90), col("number", "Number", width=130),
            col("type", "Type", width=80), col("party", "Party", width=220),
            col("method", "Method", width=100), col("ref", "Ref", width=100),
            col("received", "Received", "money"), col("paid", "Paid", "money")]
    tot = _sum_totals(rows, ("received", "paid"), "party", "TOTAL")
    return _result("Receipts & Payments", cols, rows, tot,
                   [("Received", tot["received"], "money"), ("Paid", tot["paid"], "money"),
                    ("Net", tot["received"] - tot["paid"], "money")])


# ════════════════════════════════════════════════════════════
#  SALES & POS
# ════════════════════════════════════════════════════════════
def _channel(h):
    ch = str(h.get("sale_channel") or "").upper()
    return {"POS": "POS", "INV": "Invoice"}.get(ch, "—")


@report("sales_register", GROUP_SALES, "Sales Register",
        params=("period", "party", "channel"),
        desc="Every confirmed sales invoice / POS bill.")
def r_sales_register(date_from, date_to, party="", channel="All", **_):
    rows = []
    ch = (channel or "All").lower()
    for h in _sales_headers(date_from, date_to):
        if not _match(party, h.get("customer_name_en"), h.get("invoice_number")):
            continue
        c = _channel(h)
        if ch == "pos" and c != "POS":
            continue
        if ch.startswith("inv") and c == "POS":
            continue
        sg = -1 if _is_credit_note(h) else 1
        tot = sg * _f(h.get("total_amount"))
        paid = sg * min(_f(h.get("paid_amount")), _f(h.get("total_amount")))
        rows.append({"date": _d(h["invoice_date"]), "number": h["invoice_number"],
                     "type": _doc_kind(h),
                     "customer": h.get("customer_name_en") or "",
                     "vat_no": h.get("customer_vat_no") or "", "channel": c,
                     "method": (h.get("payment_method") or "").replace("_", " ").title(),
                     "discount": round(sg * _f(h.get("discount_amount")), 2),
                     "net": round(sg * _f(h.get("taxable_amount")), 2),
                     "vat": round(sg * _f(h.get("vat_amount")), 2), "total": round(tot, 2),
                     "paid": round(paid, 2), "balance": round(tot - paid, 2)})
    cols = [col("date", "Date", width=90), col("number", "Invoice #", width=100),
            col("type", "Type", width=118),
            col("customer", "Customer", width=200), col("vat_no", "Customer VAT No.", width=140),
            col("channel", "Channel", width=70),
            col("method", "Payment", width=80), col("discount", "Discount", "money", 85),
            col("net", "Net", "money"),
            col("vat", "VAT", "money"), col("total", "Total", "money"),
            col("paid", "Paid", "money"), col("balance", "Balance", "money")]
    tot = _sum_totals(rows, ("discount", "net", "vat", "total", "paid", "balance"),
                      "customer", "TOTAL")
    n = len(rows)
    return _result("Sales Register", cols, rows, tot,
                   [("Bills", n, "int"), ("Net sales", tot["net"], "money"),
                    ("VAT", tot["vat"], "money"), ("Total", tot["total"], "money"),
                    ("Average bill", tot["total"] / n if n else 0, "money")],
                   ["'Paid' = money taken on the bill + receipts settled against it. "
                    "Advances not linked to a bill show in Customer Balances."])


def _item_sales(f, t):
    fm = _factor_map(); prods = _products(); cats = _categories()
    agg = OrderedDict()
    for ln in _sales_lines(f, t):
        pid = ln.get("product_id")
        sg = -1 if str(ln.get("invoice_type_code") or "388") == "381" else 1
        k = pid or ("x", ln.get("product_name_en"))
        a = agg.setdefault(k, {"pid": pid, "code": ln.get("product_code") or "",
                               "name": ln.get("product_name_en") or "", "qty": 0.0,
                               "net": 0.0, "vat": 0.0, "total": 0.0, "r_qty": 0.0,
                               "r_net": 0.0, "bills": set()})
        a["qty"] += sg * _f(ln["quantity"]) * _factor(fm, pid, ln.get("unit_of_measure_en"))
        a["net"] += sg * _f(ln["line_amount"]); a["vat"] += sg * _f(ln["vat_amount"])
        a["total"] += sg * _f(ln["line_total"]); a["bills"].add(ln["invoice_id"])
    for ln in _return_lines("sales", f, t):
        pid = ln.get("product_id")
        k = pid or ("x", ln.get("product_name_en"))
        a = agg.setdefault(k, {"pid": pid, "code": ln.get("product_code") or "",
                               "name": ln.get("product_name_en") or "", "qty": 0.0,
                               "net": 0.0, "vat": 0.0, "total": 0.0, "r_qty": 0.0,
                               "r_net": 0.0, "bills": set()})
        a["r_qty"] += _f(ln["quantity"]) * _factor(fm, pid, ln.get("unit_of_measure_en"))
        a["r_net"] += _f(ln["line_amount"])
    for a in agg.values():
        p = prods.get(a["pid"]) or {}
        a["category"] = cats.get(p.get("category_id"), "") or "(no category)"
        c = _cost(p) if (str(p.get("product_type") or "product").lower() != "service") else 0.0
        a["cost"] = round((a["qty"] - a["r_qty"]) * c, 2)
        a["net_sales"] = round(a["net"] - a["r_net"], 2)
        a["profit"] = round(a["net_sales"] - a["cost"], 2)
        a["margin"] = round(100.0 * a["profit"] / a["net_sales"], 1) if a["net_sales"] else 0.0
    return list(agg.values())


@report("sales_items", GROUP_SALES, "Item-wise Sales & Profit", params=("period", "product", "category"),
        desc="Quantity, value, returns, cost and profit per product.")
def r_sales_items(date_from, date_to, product="", category="", **_):
    rows = []
    for a in _item_sales(date_from, date_to):
        if not _match(product, a["code"], a["name"]):
            continue
        if category and category != "All" and a["category"] != category:
            continue
        rows.append({"code": a["code"], "name": a["name"], "category": a["category"],
                     "qty": round(a["qty"], 3), "net": round(a["net"], 2),
                     "r_qty": round(a["r_qty"], 3), "r_net": round(a["r_net"], 2),
                     "net_sales": a["net_sales"], "cost": a["cost"], "profit": a["profit"],
                     "margin": a["margin"]})
    rows.sort(key=lambda r: -r["net_sales"])
    cols = [col("code", "Code", width=80), col("name", "Product", width=200),
            col("category", "Category", width=110), col("qty", "Qty sold", "qty", 80),
            col("net", "Sales", "money"), col("r_qty", "Qty ret.", "qty", 70),
            col("r_net", "Returns", "money"), col("net_sales", "Net sales", "money"),
            col("cost", "Cost", "money"), col("profit", "Profit", "money"),
            col("margin", "Margin %", "pct", 70)]
    tot = _sum_totals(rows, ("qty", "net", "r_qty", "r_net", "net_sales", "cost", "profit"),
                      "name", "TOTAL")
    tot["margin"] = round(100.0 * tot["profit"] / tot["net_sales"], 1) if tot["net_sales"] else 0.0
    return _result("Item-wise Sales & Profit", cols, rows, tot,
                   [("Items", len(rows), "int"), ("Net sales", tot["net_sales"], "money"),
                    ("Cost", tot["cost"], "money"), ("Gross profit", tot["profit"], "money"),
                    ("Margin", tot["margin"], "pct")],
                   ["Cost uses each product's current average cost."])


@report("sales_category", GROUP_SALES, "Category-wise Sales",
        desc="Sales, cost and profit per product category.")
def r_sales_category(date_from, date_to, **_):
    by = OrderedDict()
    for a in _item_sales(date_from, date_to):
        c = by.setdefault(a["category"], {"category": a["category"], "items": 0, "qty": 0.0,
                                          "net_sales": 0.0, "cost": 0.0, "profit": 0.0})
        c["items"] += 1; c["qty"] += a["qty"] - a["r_qty"]
        c["net_sales"] += a["net_sales"]; c["cost"] += a["cost"]; c["profit"] += a["profit"]
    rows = []
    for c in sorted(by.values(), key=lambda x: -x["net_sales"]):
        c = {k: (round(v, 2) if isinstance(v, float) else v) for k, v in c.items()}
        c["margin"] = round(100.0 * c["profit"] / c["net_sales"], 1) if c["net_sales"] else 0.0
        rows.append(c)
    tot_ns = sum(r["net_sales"] for r in rows)
    for r in rows:
        r["share"] = round(100.0 * r["net_sales"] / tot_ns, 1) if tot_ns else 0.0
    cols = [col("category", "Category", width=200), col("items", "Items", "int", 60),
            col("qty", "Qty", "qty", 80), col("net_sales", "Net sales", "money"),
            col("share", "Share %", "pct", 70), col("cost", "Cost", "money"),
            col("profit", "Profit", "money"), col("margin", "Margin %", "pct", 70)]
    tot = _sum_totals(rows, ("items", "qty", "net_sales", "cost", "profit"), "category", "TOTAL")
    return _result("Category-wise Sales", cols, rows, tot)


@report("sales_customers", GROUP_SALES, "Customer-wise Sales", params=("period", "party"),
        desc="Bills, sales, returns and balance per customer.")
def r_sales_customers(date_from, date_to, party="", **_):
    by = OrderedDict()
    for h in _sales_headers(date_from, date_to):
        name = h.get("customer_name_en") or "(no name)"
        if not _match(party, name):
            continue
        sg = -1 if _is_credit_note(h) else 1
        c = by.setdefault(h.get("customer_id") or name, {"customer": name, "bills": 0,
                          "net": 0.0, "vat": 0.0, "total": 0.0, "returns": 0.0, "unpaid": 0.0})
        c["bills"] += 1; c["net"] += sg * _f(h.get("taxable_amount"))
        c["vat"] += sg * _f(h.get("vat_amount")); c["total"] += sg * _f(h.get("total_amount"))
        c["unpaid"] += sg * (_f(h.get("total_amount")) - min(_f(h.get("paid_amount")),
                                                            _f(h.get("total_amount"))))
    for r in _return_headers("sales", date_from, date_to):
        key = r.get("customer_id") or (r.get("customer_name_en") or "(no name)")
        if key in by:
            by[key]["returns"] += _f(r.get("total_amount"))
        elif _match(party, r.get("customer_name_en")):
            by[key] = {"customer": r.get("customer_name_en") or "", "bills": 0, "net": 0.0,
                       "vat": 0.0, "total": 0.0, "returns": _f(r.get("total_amount")),
                       "unpaid": 0.0}
    rows = []
    for c in by.values():
        c["net_total"] = c["total"] - c["returns"]
        rows.append({k: (round(v, 2) if isinstance(v, float) else v) for k, v in c.items()})
    rows.sort(key=lambda r: -r["net_total"])
    cols = [col("customer", "Customer", width=220), col("bills", "Bills", "int", 60),
            col("net", "Net", "money"), col("vat", "VAT", "money"), col("total", "Total", "money"),
            col("returns", "Returns", "money"), col("net_total", "Net total", "money"),
            col("unpaid", "Unpaid bills", "money")]
    tot = _sum_totals(rows, ("bills", "net", "vat", "total", "returns", "net_total", "unpaid"),
                      "customer", "TOTAL")
    return _result("Customer-wise Sales", cols, rows, tot)


@report("loyalty_points", GROUP_SALES, "⭐ Loyalty Points", params=("period", "party"),
        desc="Points each customer earned and spent in the period, and has now.")
def r_loyalty_points(date_from, date_to, party="", **_):
    from modules import loyalty
    rows = []
    for r in loyalty.summary(date_from, date_to):
        if not _match(party, r.get("customer_name_en"), r.get("customer_code"), r.get("phone")):
            continue
        bal = int(_f(r.get("balance")))
        rows.append({"code": r.get("customer_code") or "", "customer": r.get("customer_name_en") or "",
                     "phone": r.get("phone") or "", "earned": int(_f(r.get("earned"))),
                     "spent": int(_f(r.get("spent"))), "balance": bal,
                     "value": loyalty.value_of(bal) if bal > 0 else 0.0})
    cols = [col("code", "Code", width=80), col("customer", "Customer", width=220),
            col("phone", "Mobile", width=110), col("earned", "Earned", "int", 80),
            col("spent", "Spent", "int", 80), col("balance", "Points now", "int", 90),
            col("value", "Worth now (SAR)", "money")]
    tot = _sum_totals(rows, ("earned", "spent", "balance", "value"), "customer", "TOTAL")
    notes = ["1 point per %.2f SAR spent · 1 point = %.2f SAR · loyalty is %s (Settings ▸ Customer)"
             % (loyalty.sar_per_point(), loyalty.point_value(),
                "ON" if loyalty.enabled() else "OFF")]
    return _result("Loyalty Points", cols, rows, tot, notes=notes)


@report("sales_daily", GROUP_SALES, "Daily Sales Summary",
        desc="Day by day: bills, cash / card / credit, VAT, returns.")
def r_sales_daily(date_from, date_to, **_):
    by = OrderedDict()

    def day(d):
        return by.setdefault(d, {"date": d, "bills": 0, "cash": 0.0, "card": 0.0,
                                 "credit": 0.0, "net": 0.0, "vat": 0.0, "total": 0.0,
                                 "returns": 0.0})
    for h in _sales_headers(date_from, date_to):
        r = day(_d(h["invoice_date"]))
        sg = -1 if _is_credit_note(h) else 1
        tot = sg * _f(h.get("total_amount"))
        r["bills"] += 1; r["net"] += sg * _f(h.get("taxable_amount"))
        r["vat"] += sg * _f(h.get("vat_amount")); r["total"] += tot
        m = str(h.get("payment_method") or "cash").lower()
        paid = sg * min(_f(h.get("paid_amount")), _f(h.get("total_amount")))
        sp = _split(h)
        if m == "credit":
            r["credit"] += tot
        elif sp:
            r["cash"] += sg * sp[0]; r["card"] += sg * sp[1]
            r["credit"] += tot - sg * (sp[0] + sp[1])
        elif m in ("card", "bank", "bank_transfer", "mada", "cheque"):
            r["card"] += paid; r["credit"] += tot - paid
        else:
            r["cash"] += paid; r["credit"] += tot - paid
    for h in _return_headers("sales", date_from, date_to):
        day(_d(h["return_date"]))["returns"] += _f(h.get("total_amount"))
    rows = []
    for d in sorted(by):
        r = by[d]
        r["net_total"] = r["total"] - r["returns"]
        rows.append({k: (round(v, 2) if isinstance(v, float) else v) for k, v in r.items()})
    cols = [col("date", "Date", width=100), col("bills", "Bills", "int", 60),
            col("cash", "Cash", "money"), col("card", "Card / Bank", "money"),
            col("credit", "Credit", "money"), col("net", "Net", "money"),
            col("vat", "VAT", "money"), col("total", "Total", "money"),
            col("returns", "Returns", "money"), col("net_total", "Net total", "money")]
    tot = _sum_totals(rows, ("bills", "cash", "card", "credit", "net", "vat", "total",
                             "returns", "net_total"), "date", "TOTAL")
    return _result("Daily Sales Summary", cols, rows, tot,
                   [("Days", len(rows), "int"), ("Total sales", tot["total"], "money"),
                    ("Returns", tot["returns"], "money"), ("Net", tot["net_total"], "money")])


@report("bill_profit", GROUP_SALES, "Bill-wise Profit",
        desc="Net amount, cost and profit of every bill.")
def r_bill_profit(date_from, date_to, **_):
    fm = _factor_map(); prods = _products()
    cost = defaultdict(float)
    for ln in _sales_lines(date_from, date_to):
        p = prods.get(ln.get("product_id")) or {}
        if str(p.get("product_type") or "product").lower() == "service":
            continue
        cost[ln["invoice_id"]] += _f(ln["quantity"]) * _factor(
            fm, ln.get("product_id"), ln.get("unit_of_measure_en")) * _cost(p)
    rows = []
    for h in _sales_headers(date_from, date_to):
        sg = -1 if _is_credit_note(h) else 1
        net = sg * _f(h.get("taxable_amount")); c = sg * cost[h["invoice_id"]]
        rows.append({"date": _d(h["invoice_date"]), "number": h["invoice_number"],
                     "customer": h.get("customer_name_en") or "", "channel": _channel(h),
                     "net": round(net, 2), "cost": round(c, 2), "profit": round(net - c, 2),
                     "margin": round(100.0 * (net - c) / net, 1) if net else 0.0})
    cols = [col("date", "Date", width=90), col("number", "Bill #", width=100),
            col("customer", "Customer", width=200), col("channel", "Channel", width=70),
            col("net", "Net", "money"), col("cost", "Cost", "money"),
            col("profit", "Profit", "money"), col("margin", "Margin %", "pct", 70)]
    tot = _sum_totals(rows, ("net", "cost", "profit"), "customer", "TOTAL")
    tot["margin"] = round(100.0 * tot["profit"] / tot["net"], 1) if tot["net"] else 0.0
    return _result("Bill-wise Profit", cols, rows, tot,
                   [("Net sales", tot["net"], "money"), ("Cost", tot["cost"], "money"),
                    ("Profit", tot["profit"], "money"), ("Margin", tot["margin"], "pct")],
                   ["Per bill, before returns — see Item-wise Sales & Profit for "
                    "profit after returns."])


def _pos_headers(date_from, date_to):
    heads = _sales_headers(date_from, date_to)
    if not _has_col("sales_invoices", "sale_channel"):
        return heads, False
    pos = [h for h in heads if str(h.get("sale_channel") or "").upper() == "POS"]
    return pos, True


@report("pos_bills", GROUP_SALES, "POS Bills", params=("period", "method"),
        desc="POS bills with cashier, payment type and totals.")
def r_pos_bills(date_from, date_to, method="All", **_):
    heads, tagged = _pos_headers(date_from, date_to)
    users = _users()
    rows = []
    m = (method or "All").lower()
    by = defaultdict(float)
    for h in heads:
        pm = str(h.get("payment_method") or "cash").lower()
        if m not in ("all", "") and pm != m:
            continue
        by[pm] += _f(h.get("total_amount"))
        tm = h.get("invoice_time")
        rows.append({"date": _d(h["invoice_date"]),
                     "time": str(tm)[:5] if tm else "",
                     "number": h["invoice_number"],
                     "cashier": users.get(h.get("created_by"), "") or "",
                     "customer": h.get("customer_name_en") or "",
                     "method": pm.replace("_", " ").title(),
                     "net": round(_f(h.get("taxable_amount")), 2),
                     "vat": round(_f(h.get("vat_amount")), 2),
                     "total": round(_f(h.get("total_amount")), 2)})
    cols = [col("date", "Date", width=90), col("time", "Time", width=60),
            col("number", "Bill #", width=90), col("cashier", "Cashier", width=120),
            col("customer", "Customer", width=180), col("method", "Payment", width=80),
            col("net", "Net", "money"), col("vat", "VAT", "money"), col("total", "Total", "money")]
    tot = _sum_totals(rows, ("net", "vat", "total"), "customer", "TOTAL")
    n = len(rows)
    summ = [("Bills", n, "int"), ("Total", tot["total"], "money"),
            ("Average bill", tot["total"] / n if n else 0, "money")]
    summ += [(k.replace("_", " ").title(), v, "money") for k, v in sorted(by.items())]
    notes = [] if tagged else ["POS bills are tagged from v5.7 on; older bills show under "
                               "Sales Register."]
    return _result("POS Bills", cols, rows, tot, summ, notes)


@report("pos_summary", GROUP_SALES, "POS Summary",
        desc="POS takings per day split by cash, card and credit, per cashier.")
def r_pos_summary(date_from, date_to, **_):
    heads, tagged = _pos_headers(date_from, date_to)
    users = _users()
    by = OrderedDict(); cashiers = defaultdict(float)
    for h in heads:
        d = _d(h["invoice_date"])
        r = by.setdefault(d, {"date": d, "bills": 0, "cash": 0.0, "card": 0.0,
                              "credit": 0.0, "vat": 0.0, "total": 0.0})
        pm = str(h.get("payment_method") or "cash").lower()
        tot = _f(h.get("total_amount"))
        r["bills"] += 1; r["vat"] += _f(h.get("vat_amount")); r["total"] += tot
        sp = _split(h)
        if pm == "credit":
            r["credit"] += tot
        elif sp:
            r["cash"] += sp[0]; r["card"] += sp[1]
            r["credit"] += max(0.0, tot - sp[0] - sp[1])
        elif pm == "cash":
            r["cash"] += tot
        else:
            r["card"] += tot
        cashiers[users.get(h.get("created_by"), "") or "(not recorded)"] += tot
    rows = []
    for d in sorted(by):
        r = by[d]
        r["avg"] = r["total"] / r["bills"] if r["bills"] else 0.0
        rows.append({k: (round(v, 2) if isinstance(v, float) else v) for k, v in r.items()})
    cols = [col("date", "Date", width=100), col("bills", "Bills", "int", 60),
            col("cash", "Cash", "money"), col("card", "Card / Bank", "money"),
            col("credit", "Credit", "money"), col("vat", "VAT", "money"),
            col("total", "Total", "money"), col("avg", "Avg bill", "money")]
    tot = _sum_totals(rows, ("bills", "cash", "card", "credit", "vat", "total"), "date", "TOTAL")
    tot["avg"] = round(tot["total"] / tot["bills"], 2) if tot["bills"] else 0.0
    summ = [("Bills", tot["bills"], "int"), ("Cash", tot["cash"], "money"),
            ("Card / Bank", tot["card"], "money"), ("Credit", tot["credit"], "money")]
    summ += [("Cashier: " + k, v, "money") for k, v in sorted(cashiers.items(), key=lambda x: -x[1])[:4]]
    notes = [] if tagged else ["POS bills are tagged from v5.7 on."]
    return _result("POS Summary", cols, rows, tot, summ, notes)


@report("counter_sales", GROUP_SALES, "Counter & Branch-wise Sales",
        desc="Takings per counter (till) and per branch: bills, cash, card, credit, "
             "returns and the cash that should be in each drawer.")
def r_counter_sales(date_from, date_to, **_):
    try:
        from modules import branches as BR
        data = BR.takings(date_from, date_to)
        counters = {c["counter_id"]: c for c in BR.list_counters()}
        branches = {b["branch_id"]: b["name_en"] for b in BR.list_branches()}
    except Exception:
        data, counters, branches = {}, {}, {}
    by_branch = OrderedDict()
    for cid, x in data.items():
        c = counters.get(cid) or {}
        bname = branches.get(c.get("branch_id"), "(not recorded)") if c else "(not recorded)"
        by_branch.setdefault(bname, []).append((c.get("name_en") or "(not recorded)", x))
    keys = ("bills", "total", "cash", "card", "credit", "returns", "receipts",
            "paid_out", "expected_cash")
    rows = []
    for bname in sorted(by_branch):
        rows.append({"counter": bname, "_style": "head"})
        sub = {k: 0.0 for k in keys}
        for cname, x in sorted(by_branch[bname]):
            r = {"counter": "    " + cname}
            for k in keys:
                r[k] = x.get(k, 0)
                sub[k] += _f(x.get(k))
            rows.append(r)
        rows.append(dict({"counter": "  Total " + bname, "_style": "total"},
                         **{k: round(v, 2) for k, v in sub.items()}))
    tot = {k: round(sum(_f(x.get(k)) for x in data.values()), 2) for k in keys}
    tot["counter"] = "TOTAL"
    cols = [col("counter", "Branch / Counter", width=200), col("bills", "Bills", "int", 60),
            col("total", "Sales", "money"), col("cash", "Cash", "money"),
            col("card", "Card / Bank", "money"), col("credit", "Credit", "money"),
            col("returns", "Returns", "money"), col("receipts", "Cash receipts", "money"),
            col("paid_out", "Cash paid out", "money"),
            col("expected_cash", "Cash in drawer", "money")]
    summ = [("Bills", tot["bills"], "int"), ("Sales", tot["total"], "money"),
            ("Cash", tot["cash"], "money"), ("Card / Bank", tot["card"], "money")]
    return _result("Counter & Branch-wise Sales", cols, rows, tot, summ,
                   ["Cash in drawer = cash on bills + cash receipts − cash refunds − "
                    "cash paid out at that counter (expenses, supplier payments and "
                    "cash purchases)."])


@report("sales_returns", GROUP_SALES, "Sales Return Register", params=("period", "party"),
        desc="Confirmed sales returns / credit notes.")
def r_sales_returns(date_from, date_to, party="", **_):
    rows = []
    for h in _return_headers("sales", date_from, date_to):
        if not _match(party, h.get("customer_name_en"), h.get("return_number")):
            continue
        rows.append({"date": _d(h["return_date"]), "number": h["return_number"],
                     "orig": h.get("original_invoice_no") or "",
                     "customer": h.get("customer_name_en") or "",
                     "vat_no": h.get("customer_vat_no") or "",
                     "method": (h.get("payment_method") or "").replace("_", " ").title(),
                     "discount": round(_f(h.get("discount_amount")), 2),
                     "net": round(_f(h.get("taxable_amount")), 2),
                     "vat": round(_f(h.get("vat_amount")), 2),
                     "total": round(_f(h.get("total_amount")), 2)})
    cols = [col("date", "Date", width=90), col("number", "Return #", width=110),
            col("orig", "Orig. invoice", width=100), col("customer", "Customer", width=200),
            col("vat_no", "Customer VAT No.", width=140),
            col("method", "Refund", width=80), col("discount", "Discount", "money", 85),
            col("net", "Net", "money"),
            col("vat", "VAT", "money"), col("total", "Total", "money")]
    tot = _sum_totals(rows, ("discount", "net", "vat", "total"), "customer", "TOTAL")
    return _result("Sales Return Register", cols, rows, tot,
                   [("Returns", len(rows), "int"), ("Total", tot["total"], "money"),
                    ("VAT reversed", tot["vat"], "money")])


@report("sales_lines", GROUP_SALES, "Sales Register — item details",
        params=("period", "party", "product"),
        desc="Every line of every confirmed sales invoice / POS bill: product, qty, "
             "rate, discount, net, VAT and total. Credit notes are minus.")
def r_sales_lines(date_from, date_to, party="", product="", **_):
    rows = []
    for ln in _q("SELECT sil.*, si.invoice_date, si.invoice_number, si.invoice_type_code, "
                 "si.customer_name_en AS cust_name, si.customer_vat_no AS cust_vat "
                 "FROM sales_invoice_lines sil JOIN sales_invoices si ON si.invoice_id=sil.invoice_id "
                 "WHERE si.invoice_status='confirmed' AND si.invoice_date BETWEEN ? AND ? "
                 "ORDER BY si.invoice_date, si.invoice_id, sil.line_number, sil.line_id",
                 (date_from, date_to)):
        if not _match(party, ln.get("cust_name"), ln.get("invoice_number")):
            continue
        if not _match(product, ln.get("product_code"), ln.get("product_name_en"),
                      ln.get("product_name_ar")):
            continue
        sg = -1 if str(ln.get("invoice_type_code") or "388") == "381" else 1
        net = sg * _f(ln.get("line_amount"))
        vat = sg * _f(ln.get("vat_amount"))
        rows.append({"date": _d(ln["invoice_date"]), "number": ln["invoice_number"],
                     "customer": ln.get("cust_name") or "", "code": ln.get("product_code") or "",
                     "product": ln.get("product_name_en") or "",
                     "product_ar": ln.get("product_name_ar") or "",
                     "qty": round(sg * _f(ln.get("quantity")), 3),
                     "unit": ln.get("unit_of_measure_en") or "",
                     "rate": round(_f(ln.get("unit_price")), 2),
                     "discount": round(sg * _f(ln.get("discount_amount")), 2),
                     "net": round(net, 2), "vat_rate": _f(ln.get("vat_rate")),
                     "vat": round(vat, 2), "total": round(net + vat, 2)})
    cols = [col("date", "Date", width=90), col("number", "Invoice #", width=100),
            col("customer", "Customer", width=170), col("code", "Code", width=80),
            col("product", "Product", width=200), col("product_ar", "Product (Arabic)", width=170),
            col("qty", "Qty", "qty", 70), col("unit", "Unit", width=60),
            col("rate", "Rate", "money", 85), col("discount", "Discount", "money", 85),
            col("net", "Net", "money"), col("vat_rate", "VAT %", "pct", 60),
            col("vat", "VAT", "money"), col("total", "Total", "money")]
    tot = _sum_totals(rows, ("qty", "discount", "net", "vat", "total"), "product", "TOTAL")
    return _result("Sales Register — item details", cols, rows, tot,
                   [("Lines", len(rows), "int"), ("Net", tot["net"], "money"),
                    ("VAT", tot["vat"], "money"), ("Total", tot["total"], "money")])


# ════════════════════════════════════════════════════════════
#  PURCHASES
# ════════════════════════════════════════════════════════════
@report("purchase_register", GROUP_PUR, "Purchase Register", params=("period", "party"),
        desc="Every confirmed purchase invoice.")
def r_purchase_register(date_from, date_to, party="", **_):
    rows = []
    for h in _purchase_headers(date_from, date_to):
        if not _match(party, h.get("supplier_name_en"), h.get("invoice_number"),
                      h.get("supplier_invoice_no")):
            continue
        tot = _f(h.get("total_amount"))
        paid = min(_f(h.get("paid_amount")), tot)
        rows.append({"date": _d(h["invoice_date"]), "number": h["invoice_number"],
                     "sup_inv": h.get("supplier_invoice_no") or "",
                     "supplier": h.get("supplier_name_en") or "",
                     "vat_no": h.get("supplier_vat_no") or "",
                     "method": (h.get("payment_method") or "credit").replace("_", " ").title(),
                     "discount": round(_f(h.get("discount_amount")), 2),
                     "net": round(_f(h.get("taxable_amount")), 2),
                     "vat": round(_f(h.get("vat_amount")), 2), "total": round(tot, 2),
                     "balance": round(tot - paid, 2)})
    cols = [col("date", "Date", width=90), col("number", "Purchase #", width=100),
            col("sup_inv", "Supplier inv.", width=100), col("supplier", "Supplier", width=200),
            col("vat_no", "Supplier VAT No.", width=140),
            col("method", "Payment", width=80), col("discount", "Discount", "money", 85),
            col("net", "Net", "money"),
            col("vat", "VAT", "money"), col("total", "Total", "money"),
            col("balance", "Unpaid", "money")]
    tot = _sum_totals(rows, ("discount", "net", "vat", "total", "balance"), "supplier", "TOTAL")
    return _result("Purchase Register", cols, rows, tot,
                   [("Invoices", len(rows), "int"), ("Net", tot["net"], "money"),
                    ("Input VAT", tot["vat"], "money"), ("Total", tot["total"], "money")],
                   ["'Unpaid' is per bill after supplier payments settled against it. "
                    "Advances not linked to a bill show in Supplier Balances."])


@report("purchase_items", GROUP_PUR, "Item-wise Purchases", params=("period", "product"),
        desc="Quantity, value, returns and average rate per product.")
def r_purchase_items(date_from, date_to, product="", **_):
    fm = _factor_map()
    agg = OrderedDict()
    for ln in _purchase_lines(date_from, date_to):
        pid = ln.get("product_id")
        a = agg.setdefault(pid or ln.get("product_name_en"), {
            "code": ln.get("product_code") or "", "name": ln.get("product_name_en") or "",
            "qty": 0.0, "net": 0.0, "vat": 0.0, "r_qty": 0.0, "r_net": 0.0})
        a["qty"] += _f(ln["quantity"]) * _factor(fm, pid, ln.get("unit_of_measure_en"))
        a["net"] += _f(ln["line_amount"]); a["vat"] += _f(ln["vat_amount"])
    for ln in _return_lines("purchase", date_from, date_to):
        pid = ln.get("product_id")
        a = agg.setdefault(pid or ln.get("product_name_en"), {
            "code": ln.get("product_code") or "", "name": ln.get("product_name_en") or "",
            "qty": 0.0, "net": 0.0, "vat": 0.0, "r_qty": 0.0, "r_net": 0.0})
        a["r_qty"] += _f(ln["quantity"]) * _factor(fm, pid, ln.get("unit_of_measure_en"))
        a["r_net"] += _f(ln["line_amount"])
    rows = []
    for a in agg.values():
        if not _match(product, a["code"], a["name"]):
            continue
        a["net_value"] = a["net"] - a["r_net"]
        a["rate"] = a["net"] / a["qty"] if a["qty"] else 0.0
        rows.append({k: (round(v, 3 if k in ("qty", "r_qty") else 2) if isinstance(v, float) else v)
                     for k, v in a.items()})
    rows.sort(key=lambda r: -r["net_value"])
    cols = [col("code", "Code", width=80), col("name", "Product", width=220),
            col("qty", "Qty (base)", "qty", 90), col("rate", "Avg cost", "money"),
            col("net", "Value", "money"), col("vat", "VAT", "money"),
            col("r_qty", "Qty ret.", "qty", 80), col("r_net", "Returns", "money"),
            col("net_value", "Net value", "money")]
    tot = _sum_totals(rows, ("qty", "net", "vat", "r_qty", "r_net", "net_value"), "name", "TOTAL")
    tot["rate"] = None
    return _result("Item-wise Purchases", cols, rows, tot)


@report("purchase_suppliers", GROUP_PUR, "Supplier-wise Purchases", params=("period", "party"),
        desc="Purchases, returns and unpaid per supplier.")
def r_purchase_suppliers(date_from, date_to, party="", **_):
    by = OrderedDict()
    for h in _purchase_headers(date_from, date_to):
        name = h.get("supplier_name_en") or "(no name)"
        if not _match(party, name):
            continue
        s = by.setdefault(h.get("supplier_id") or name, {"supplier": name, "invoices": 0,
                          "net": 0.0, "vat": 0.0, "total": 0.0, "returns": 0.0, "unpaid": 0.0})
        tot = _f(h.get("total_amount"))
        s["invoices"] += 1; s["net"] += _f(h.get("taxable_amount"))
        s["vat"] += _f(h.get("vat_amount")); s["total"] += tot
        s["unpaid"] += tot - min(_f(h.get("paid_amount")), tot)
    for r in _return_headers("purchase", date_from, date_to):
        key = r.get("supplier_id") or (r.get("supplier_name_en") or "(no name)")
        if key in by:
            by[key]["returns"] += _f(r.get("total_amount"))
    rows = []
    for s in by.values():
        s["net_total"] = s["total"] - s["returns"]
        rows.append({k: (round(v, 2) if isinstance(v, float) else v) for k, v in s.items()})
    rows.sort(key=lambda r: -r["net_total"])
    cols = [col("supplier", "Supplier", width=220), col("invoices", "Invoices", "int", 70),
            col("net", "Net", "money"), col("vat", "VAT", "money"), col("total", "Total", "money"),
            col("returns", "Returns", "money"), col("net_total", "Net total", "money"),
            col("unpaid", "Unpaid", "money")]
    tot = _sum_totals(rows, ("invoices", "net", "vat", "total", "returns", "net_total", "unpaid"),
                      "supplier", "TOTAL")
    return _result("Supplier-wise Purchases", cols, rows, tot)


@report("purchase_returns", GROUP_PUR, "Purchase Return Register", params=("period", "party"),
        desc="Confirmed purchase returns / debit notes.")
def r_purchase_returns(date_from, date_to, party="", **_):
    rows = []
    for h in _return_headers("purchase", date_from, date_to):
        if not _match(party, h.get("supplier_name_en"), h.get("return_number")):
            continue
        rows.append({"date": _d(h["return_date"]), "number": h["return_number"],
                     "orig": h.get("original_purchase_no") or "",
                     "supplier": h.get("supplier_name_en") or "",
                     "vat_no": h.get("supplier_vat_no") or "",
                     "method": (h.get("payment_method") or "").replace("_", " ").title(),
                     "discount": round(_f(h.get("discount_amount")), 2),
                     "net": round(_f(h.get("taxable_amount")), 2),
                     "vat": round(_f(h.get("vat_amount")), 2),
                     "total": round(_f(h.get("total_amount")), 2)})
    cols = [col("date", "Date", width=90), col("number", "Return #", width=110),
            col("orig", "Orig. purchase", width=110), col("supplier", "Supplier", width=200),
            col("vat_no", "Supplier VAT No.", width=140),
            col("method", "Settled by", width=90), col("discount", "Discount", "money", 85),
            col("net", "Net", "money"),
            col("vat", "VAT", "money"), col("total", "Total", "money")]
    tot = _sum_totals(rows, ("discount", "net", "vat", "total"), "supplier", "TOTAL")
    return _result("Purchase Return Register", cols, rows, tot,
                   [("Returns", len(rows), "int"), ("Total", tot["total"], "money")])


@report("purchase_lines", GROUP_PUR, "Purchase Register — item details",
        params=("period", "party", "product"),
        desc="Every line of every confirmed purchase invoice: product, qty, cost, "
             "discount, net, VAT, total, batch and expiry.")
def r_purchase_lines(date_from, date_to, party="", product="", **_):
    rows = []
    for ln in _q("SELECT pil.*, pi.invoice_date, pi.invoice_number, pi.supplier_invoice_no, "
                 "pi.supplier_name_en AS sup_name FROM purchase_invoice_lines pil "
                 "JOIN purchase_invoices pi ON pi.purchase_id=pil.purchase_id "
                 "WHERE pi.invoice_status='confirmed' AND pi.invoice_date BETWEEN ? AND ? "
                 "ORDER BY pi.invoice_date, pi.purchase_id, pil.line_number, pil.line_id",
                 (date_from, date_to)):
        if not _match(party, ln.get("sup_name"), ln.get("invoice_number"),
                      ln.get("supplier_invoice_no")):
            continue
        if not _match(product, ln.get("product_code"), ln.get("product_name_en"),
                      ln.get("product_name_ar")):
            continue
        net, vat = _f(ln.get("line_amount")), _f(ln.get("vat_amount"))
        rows.append({"date": _d(ln["invoice_date"]), "number": ln["invoice_number"],
                     "sup_inv": ln.get("supplier_invoice_no") or "",
                     "supplier": ln.get("sup_name") or "", "code": ln.get("product_code") or "",
                     "product": ln.get("product_name_en") or "",
                     "qty": round(_f(ln.get("quantity")), 3),
                     "unit": ln.get("unit_of_measure_en") or "",
                     "rate": round(_f(ln.get("unit_price")), 2),
                     "discount": round(_f(ln.get("discount_amount")), 2),
                     "net": round(net, 2), "vat_rate": _f(ln.get("vat_rate")),
                     "vat": round(vat, 2), "total": round(net + vat, 2),
                     "batch": ln.get("batch_no") or "", "expiry": _d(ln.get("expiry_date")) or ""})
    cols = [col("date", "Date", width=90), col("number", "Purchase #", width=100),
            col("sup_inv", "Supplier inv.", width=100), col("supplier", "Supplier", width=170),
            col("code", "Code", width=80), col("product", "Product", width=200),
            col("qty", "Qty", "qty", 70), col("unit", "Unit", width=60),
            col("rate", "Cost", "money", 85), col("discount", "Discount", "money", 85),
            col("net", "Net", "money"), col("vat_rate", "VAT %", "pct", 60),
            col("vat", "VAT", "money"), col("total", "Total", "money"),
            col("batch", "Batch", width=80), col("expiry", "Expiry", width=90)]
    tot = _sum_totals(rows, ("qty", "discount", "net", "vat", "total"), "product", "TOTAL")
    return _result("Purchase Register — item details", cols, rows, tot,
                   [("Lines", len(rows), "int"), ("Net", tot["net"], "money"),
                    ("Input VAT", tot["vat"], "money"), ("Total", tot["total"], "money")])


# ════════════════════════════════════════════════════════════
#  INVENTORY
# ════════════════════════════════════════════════════════════
def _stock_items():
    return [p for p in _products().values()
            if p.get("is_active", 1) in (1, True, None)
            and str(p.get("product_type") or "product").lower() != "service"]


def _status(p):
    q = _f(p.get("stock_quantity"))
    lvl = _f(p.get("reorder_level")) or _f(p.get("min_stock"))
    if q < 0:
        return "Negative"
    if q == 0:
        return "Out"
    if lvl and q <= lvl:
        return "Low"
    return "OK"


def _base_units():
    out = {}
    if not _has_table("product_units"):
        return out
    for r in _q("SELECT product_id, unit_name_en, conversion_factor FROM product_units "
                "WHERE is_active=1"):
        if abs(_f(r.get("conversion_factor")) - 1.0) < 1e-9 and r["product_id"] not in out:
            out[r["product_id"]] = r.get("unit_name_en") or ""
    return out


@report("stock", GROUP_INV, "Stock Report",
        params=("as_now", "product", "category", "stock_filter"),
        desc="Current stock, cost value and selling value per product.")
def r_stock(product="", category="", stock_filter="All", **_):
    cats = _categories(); units = _base_units()
    rows = []
    sf = (stock_filter or "All").lower()
    for p in _stock_items():
        if not _match(product, p.get("product_code"), p.get("product_name_en"),
                      p.get("product_name_ar"), p.get("barcode")):
            continue
        cat = cats.get(p.get("category_id"), "") or "(no category)"
        if category and category != "All" and cat != category:
            continue
        st = _status(p)
        if sf.startswith("in") and _f(p.get("stock_quantity")) <= 0:
            continue
        if sf.startswith("low") and st != "Low":
            continue
        if sf.startswith("out") and st != "Out":
            continue
        if sf.startswith("neg") and st != "Negative":
            continue
        q = _f(p.get("stock_quantity")); c = _cost(p); sp = _f(p.get("unit_price"))
        rows.append({"code": p.get("product_code") or "", "name": p.get("product_name_en") or "",
                     "category": cat, "unit": units.get(p["product_id"]) or
                     (p.get("unit_of_measure_en") or ""),
                     "qty": round(q, 3), "cost": round(c, 2), "value": round(q * c, 2),
                     "price": round(sp, 2), "sale_value": round(q * sp, 2),
                     "reorder": _f(p.get("reorder_level")) or None, "status": st})
    rows.sort(key=lambda r: (r["category"], r["name"]))
    cols = [col("code", "Code", width=80), col("name", "Product", width=210),
            col("category", "Category", width=110), col("unit", "Unit", width=60),
            col("qty", "Stock", "qty", 80), col("cost", "Cost", "money", 80),
            col("value", "Stock value", "money"), col("price", "Sale price", "money", 80),
            col("sale_value", "Sale value", "money"), col("reorder", "Reorder", "qty", 70),
            col("status", "Status", width=70)]
    tot = _sum_totals(rows, ("qty", "value", "sale_value"), "name", "TOTAL")
    return _result("Stock Report", cols, rows, tot,
                   [("Items", len(rows), "int"), ("Stock value (cost)", tot["value"], "money"),
                    ("Sale value", tot["sale_value"], "money"),
                    ("Low / out", sum(1 for r in rows if r["status"] in ("Low", "Out", "Negative")),
                     "int")], subtitle="Current stock")


@report("stock_valuation", GROUP_INV, "Stock Valuation", params=("as_now",),
        desc="Stock value at cost and at selling price per category.")
def r_stock_valuation(**_):
    cats = _categories()
    by = OrderedDict()
    for p in _stock_items():
        q = _f(p.get("stock_quantity"))
        cat = cats.get(p.get("category_id"), "") or "(no category)"
        r = by.setdefault(cat, {"category": cat, "items": 0, "qty": 0.0, "value": 0.0,
                                "sale_value": 0.0})
        r["items"] += 1; r["qty"] += q
        r["value"] += q * _cost(p); r["sale_value"] += q * _f(p.get("unit_price"))
    rows = []
    for r in sorted(by.values(), key=lambda x: -x["value"]):
        r["margin"] = r["sale_value"] - r["value"]
        rows.append({k: (round(v, 2) if isinstance(v, float) else v) for k, v in r.items()})
    cols = [col("category", "Category", width=220), col("items", "Items", "int", 60),
            col("qty", "Qty", "qty", 90), col("value", "Value at cost", "money"),
            col("sale_value", "Value at sale price", "money"),
            col("margin", "Potential margin", "money")]
    tot = _sum_totals(rows, ("items", "qty", "value", "sale_value", "margin"), "category", "TOTAL")
    notes = []
    try:
        from modules import account_defaults as ad
        acc = _q1("SELECT ISNULL(SUM(jl.debit),0) - ISNULL(SUM(jl.credit),0) AS b FROM journal_lines jl "
                  "JOIN journal_entries je ON je.journal_id=jl.journal_id WHERE je.is_posted=1 "
                  "AND jl.account_id=?", (ad.inventory(),))
        if acc is not None:
            b = _f(acc.get("b"))
            notes.append(f"Inventory account balance in the books: {b:,.2f} "
                         f"(difference {tot['value'] - b:,.2f}).")
    except Exception:
        pass
    return _result("Stock Valuation", cols, rows, tot,
                   [("Stock value (cost)", tot["value"], "money"),
                    ("At sale price", tot["sale_value"], "money")], notes,
                   subtitle="Current stock")


def _movements(date_from, date_to=None, product_id=None):
    """Stock movements from confirmed documents + manual adjustments,
    base units. Returns list of {pid, date, type, ref, party, qin, qout}."""
    fm = _factor_map()
    out = []
    t = date_to or "9999-12-31"
    pf = " AND l.product_id=?" if product_id else ""
    pp = (product_id,) if product_id else ()

    def add(rows, typ, sign, dkey, rkey, party_key):
        for r in rows:
            if not r.get("product_id"):
                continue
            q = _f(r["quantity"]) * _factor(fm, r["product_id"], r.get("unit_of_measure_en"))
            if str(r.get("invoice_type_code") or "388") == "381":
                sign_ = -sign
            else:
                sign_ = sign
            out.append({"pid": r["product_id"], "date": _d(r[dkey]), "type": typ,
                        "ref": r.get(rkey) or "", "party": r.get(party_key) or "",
                        "qin": q if sign_ > 0 else 0.0, "qout": q if sign_ < 0 else 0.0})
    add(_q("SELECT l.product_id, l.quantity, l.unit_of_measure_en, h.invoice_date, "
           "h.invoice_number, h.supplier_name_en FROM purchase_invoice_lines l "
           "JOIN purchase_invoices h ON h.purchase_id=l.purchase_id "
           "WHERE h.invoice_status='confirmed' AND h.invoice_date BETWEEN ? AND ?" + pf,
           (date_from, t) + pp), "Purchase", 1, "invoice_date", "invoice_number",
        "supplier_name_en")
    add(_q("SELECT l.product_id, l.quantity, l.unit_of_measure_en, h.invoice_date, "
           "h.invoice_number, h.customer_name_en, h.invoice_type_code FROM sales_invoice_lines l "
           "JOIN sales_invoices h ON h.invoice_id=l.invoice_id "
           "WHERE h.invoice_status='confirmed' AND h.invoice_date BETWEEN ? AND ?" + pf,
           (date_from, t) + pp), "Sale", -1, "invoice_date", "invoice_number",
        "customer_name_en")
    for kind, typ, sign, pk in (("sales", "Sales return", 1, "customer_name_en"),
                                ("purchase", "Purchase return", -1, "supplier_name_en")):
        tbl = "sales_returns" if kind == "sales" else "purchase_returns"
        lt = "sales_return_lines" if kind == "sales" else "purchase_return_lines"
        if not _has_table(tbl):
            continue
        add(_q(f"SELECT l.product_id, l.quantity, l.unit_of_measure_en, h.return_date, "
               f"h.return_number, h.{pk} FROM {lt} l JOIN {tbl} h ON h.return_id=l.return_id "
               f"WHERE ISNULL(h.return_status,'confirmed')='confirmed' "
               f"AND h.return_date BETWEEN ? AND ?" + pf, (date_from, t) + pp),
            typ, sign, "return_date", "return_number", pk)
    if _has_table("inventory_transactions"):
        pf2 = " AND product_id=?" if product_id else ""
        for r in _q("SELECT product_id, txn_date, txn_type, reference_type, notes, "
                    "ISNULL(qty_in_base,0) AS qin, ISNULL(qty_out_base,0) AS qout "
                    "FROM inventory_transactions WHERE txn_date BETWEEN ? AND ? "
                    "AND (reference_type IS NULL OR reference_type NOT IN "
                    "('sales_invoice','purchase_invoice','sales_return','purchase_return'))" + pf2,
                    (date_from, t) + pp):
            rt = str(r.get("reference_type") or "manual")
            typ = _ADJ_TYPES.get(rt) or _ADJ_TYPES.get(str(r.get("txn_type") or ""), "Adjustment")
            out.append({"pid": r["product_id"], "date": _d(r["txn_date"]),
                        "type": typ, "ref": rt,
                        "party": r.get("notes") or "", "qin": _f(r["qin"]),
                        "qout": _f(r["qout"])})
    return out


_ADJ_TYPES = {"opening": "Opening stock", "stock_count": "Stock count",
              "stock_clear": "Count — cleared", "manual": "Adjustment",
              "mobile_count": "Stock count (mobile)", "van": "Van stock",
              "damage": "Damage / expiry", "transfer": "Transfer"}


@report("stock_movement", GROUP_INV, "Stock Movement Summary",
        params=("period", "product", "category"),
        desc="Opening, purchased, sold, returned, adjusted and closing stock per product.")
def r_stock_movement(date_from, date_to, product="", category="", **_):
    cats = _categories()
    items = {p["product_id"]: p for p in _stock_items()}
    after = defaultdict(float)
    nxt = (_dt.date.fromisoformat(date_to) + _dt.timedelta(days=1)).isoformat()
    for m in _movements(nxt):
        after[m["pid"]] += m["qin"] - m["qout"]
    per = defaultdict(lambda: defaultdict(float))
    for m in _movements(date_from, date_to):
        if m["type"] == "Opening stock":
            continue            # opening stock belongs in the Opening column (v5.26)
        k = {"Purchase": "pur", "Sale": "sold", "Sales return": "sret",
             "Purchase return": "pret"}.get(m["type"], "adj")
        if k == "adj":
            per[m["pid"]]["adj_in"] += m["qin"]; per[m["pid"]]["adj_out"] += m["qout"]
        elif k in ("pur", "sret"):
            per[m["pid"]][k] += m["qin"] - m["qout"]
        else:
            per[m["pid"]][k] += m["qout"] - m["qin"]
    rows = []
    for pid, p in items.items():
        if not _match(product, p.get("product_code"), p.get("product_name_en"), p.get("barcode")):
            continue
        cat = cats.get(p.get("category_id"), "") or "(no category)"
        if category and category != "All" and cat != category:
            continue
        m = per.get(pid, {})
        closing = _f(p.get("stock_quantity")) - after.get(pid, 0.0)
        net = (m.get("pur", 0) + m.get("sret", 0) + m.get("adj_in", 0)
               - m.get("sold", 0) - m.get("pret", 0) - m.get("adj_out", 0))
        opening = closing - net
        if not any(abs(x) > 1e-9 for x in (opening, closing, net)) and not m:
            continue
        rows.append({"code": p.get("product_code") or "", "name": p.get("product_name_en") or "",
                     "opening": round(opening, 3), "pur": round(m.get("pur", 0), 3),
                     "sret": round(m.get("sret", 0), 3), "adj_in": round(m.get("adj_in", 0), 3),
                     "sold": round(m.get("sold", 0), 3), "pret": round(m.get("pret", 0), 3),
                     "adj_out": round(m.get("adj_out", 0), 3), "closing": round(closing, 3),
                     "value": round(closing * _cost(p), 2)})
    rows.sort(key=lambda r: r["name"])
    cols = [col("code", "Code", width=80), col("name", "Product", width=200),
            col("opening", "Opening", "qty", 75), col("pur", "Purchased", "qty", 80),
            col("sret", "Sales ret.", "qty", 75), col("adj_in", "Adj. in", "qty", 65),
            col("sold", "Sold", "qty", 70), col("pret", "Purch. ret.", "qty", 80),
            col("adj_out", "Adj. out", "qty", 70), col("closing", "Closing", "qty", 75),
            col("value", "Closing value", "money")]
    tot = _sum_totals(rows, ("opening", "pur", "sret", "adj_in", "sold", "pret", "adj_out",
                             "closing", "value"), "name", "TOTAL")
    return _result("Stock Movement Summary", cols, rows, tot,
                   [("Products", len(rows), "int"), ("Closing value", tot["value"], "money")],
                   ["Quantities in base units. Closing is worked back from today's stock. "
                    "Opening stock (Accounts ▸ Opening Balances) is in the Opening column."])


def find_product(text):
    t = (text or "").strip()
    if not t:
        return None
    for p in _products().values():
        if t.lower() in (str(p.get("product_code") or "").lower(),
                         str(p.get("barcode") or "").lower()):
            return p
    if _has_table("product_units"):
        u = _q1("SELECT product_id FROM product_units WHERE barcode=?", (t,))
        if u:
            return _products().get(u["product_id"])
    for p in _products().values():
        if t.lower() in str(p.get("product_name_en") or "").lower():
            return p
    return None


@report("stock_ledger", GROUP_INV, "Item Ledger (Stock Card)", params=("period", "product"),
        desc="Every in / out of one product with running balance.")
def r_stock_ledger(date_from, date_to, product="", **_):
    cols = [col("date", "Date", width=90), col("type", "Type", width=120),
            col("ref", "Ref", width=110), col("party", "Party / note", width=220),
            col("qin", "In", "qty", 80), col("qout", "Out", "qty", 80),
            col("balance", "Balance", "qty", 90)]
    p = find_product(product)
    if not p:
        return _result("Item Ledger", cols, [], None, [],
                       ["Type a product code, barcode or name, then Run."])
    pid = p["product_id"]
    nxt = (_dt.date.fromisoformat(date_to) + _dt.timedelta(days=1)).isoformat()
    after = sum(m["qin"] - m["qout"] for m in _movements(nxt, None, pid))
    mv = sorted(_movements(date_from, date_to, pid), key=lambda m: (m["date"], m["type"]))
    closing = _f(p.get("stock_quantity")) - after
    opening = closing - sum(m["qin"] - m["qout"] for m in mv)
    bal = opening
    rows = [{"date": date_from, "type": "Opening", "ref": "", "party": "", "qin": None,
             "qout": None, "balance": round(bal, 3), "_style": "sub"}]
    for m in mv:
        bal += m["qin"] - m["qout"]
        rows.append({"date": m["date"], "type": m["type"], "ref": m["ref"], "party": m["party"],
                     "qin": round(m["qin"], 3) or None, "qout": round(m["qout"], 3) or None,
                     "balance": round(bal, 3)})
    tin = sum(m["qin"] for m in mv); tout = sum(m["qout"] for m in mv)
    rows.append({"date": date_to, "type": "Closing", "ref": "", "party": "",
                 "qin": round(tin, 3), "qout": round(tout, 3), "balance": round(bal, 3),
                 "_style": "total"})
    return _result(f"Item Ledger — {p.get('product_code')} {p.get('product_name_en')}",
                   cols, rows, None,
                   [("Opening", opening, "qty"), ("In", tin, "qty"), ("Out", tout, "qty"),
                    ("Closing", bal, "qty"), ("Value", bal * _cost(p), "money")])


@report("low_stock", GROUP_INV, "Low Stock / Reorder", params=("as_now", "category"),
        desc="Items at or below reorder level with a suggested order quantity.")
def r_low_stock(category="", **_):
    cats = _categories()
    rows = []
    for p in _stock_items():
        q = _f(p.get("stock_quantity"))
        lvl = _f(p.get("reorder_level")) or _f(p.get("min_stock"))
        if q > 0 and (not lvl or q > lvl):
            continue
        cat = cats.get(p.get("category_id"), "") or "(no category)"
        if category and category != "All" and cat != category:
            continue
        mx = _f(p.get("max_stock"))
        rq = _f(p.get("reorder_quantity"))
        sug = rq or (mx - q if mx > q else max(lvl * 2 - q, 0)) or 0
        c = _cost(p)
        rows.append({"code": p.get("product_code") or "", "name": p.get("product_name_en") or "",
                     "category": cat, "qty": round(q, 3), "level": lvl or None,
                     "suggest": round(sug, 3), "cost": round(c, 2),
                     "order_value": round(sug * c, 2), "status": _status(p)})
    rows.sort(key=lambda r: (r["qty"], r["name"]))
    cols = [col("code", "Code", width=80), col("name", "Product", width=220),
            col("category", "Category", width=120), col("qty", "Stock", "qty", 80),
            col("level", "Reorder at", "qty", 80), col("suggest", "Suggested order", "qty", 110),
            col("cost", "Cost", "money", 80), col("order_value", "Order value", "money"),
            col("status", "Status", width=70)]
    tot = _sum_totals(rows, ("suggest", "order_value"), "name", "TOTAL")
    return _result("Low Stock / Reorder", cols, rows, tot,
                   [("Items to reorder", len(rows), "int"),
                    ("Order value", tot["order_value"], "money")], subtitle="Current stock")


@report("dead_stock", GROUP_INV, "Non-moving / Dead Stock", params=("period", "category"),
        desc="Items in stock that did not sell in the period.")
def r_dead_stock(date_from, date_to, category="", **_):
    cats = _categories()
    sold = {r["product_id"] for r in _q(
        "SELECT DISTINCT sil.product_id FROM sales_invoice_lines sil JOIN sales_invoices si "
        "ON si.invoice_id=sil.invoice_id WHERE si.invoice_status='confirmed' "
        "AND si.invoice_date BETWEEN ? AND ?", (date_from, date_to))}
    last = {r["product_id"]: _d(r["d"]) for r in _q(
        "SELECT sil.product_id, MAX(si.invoice_date) AS d FROM sales_invoice_lines sil "
        "JOIN sales_invoices si ON si.invoice_id=sil.invoice_id "
        "WHERE si.invoice_status='confirmed' GROUP BY sil.product_id")}
    rows = []
    for p in _stock_items():
        q = _f(p.get("stock_quantity"))
        if q <= 0 or p["product_id"] in sold:
            continue
        cat = cats.get(p.get("category_id"), "") or "(no category)"
        if category and category != "All" and cat != category:
            continue
        rows.append({"code": p.get("product_code") or "", "name": p.get("product_name_en") or "",
                     "category": cat, "qty": round(q, 3), "value": round(q * _cost(p), 2),
                     "last_sold": last.get(p["product_id"], "never")})
    rows.sort(key=lambda r: -r["value"])
    cols = [col("code", "Code", width=80), col("name", "Product", width=220),
            col("category", "Category", width=120), col("qty", "Stock", "qty", 80),
            col("value", "Value at cost", "money"), col("last_sold", "Last sold", width=100)]
    tot = _sum_totals(rows, ("qty", "value"), "name", "TOTAL")
    return _result("Non-moving Stock", cols, rows, tot,
                   [("Items", len(rows), "int"), ("Value tied up", tot["value"], "money")])


# ════════════════════════════════════════════════════════════
#  v5.26 — CASH & BANK · DAY REPORT · AGING · EXPENSES · OPENINGS
# ════════════════════════════════════════════════════════════
def _cash_accounts():
    """Leaf cash / bank / card accounts (11xx + the default cash roles)."""
    rows = _q("SELECT account_id, account_code, account_name_en, account_type, "
              "parent_account_id FROM accounts ORDER BY account_code")
    parents = {r.get("parent_account_id") for r in rows if r.get("parent_account_id")}
    ids = set()
    try:
        from modules import account_defaults as ad
        for role in ("cash", "bank", "card", "savings", "gateway", "petty_cash"):
            a = getattr(ad, role)()
            if a:
                ids.add(a)
    except Exception:
        pass
    out = []
    for r in rows:
        code = str(r.get("account_code") or "")
        if r["account_id"] in parents or "-" in code:
            continue
        if r["account_id"] in ids or (r.get("account_type") == "asset" and code.startswith("11")):
            out.append(r)
    return out


@report("cash_bank", GROUP_ACC, "Cash & Bank Balances",
        desc="Opening, money in, money out and closing of every cash, bank and card account.")
def r_cash_bank(date_from, date_to, **_):
    accs = _cash_accounts()
    ids = [a["account_id"] for a in accs]
    rows = []
    if ids:
        # opening balances (Accounts ▸ Opening Balances) are the start, not money in
        opn = "ISNULL(je.source_type,'') LIKE 'opening%%'"     # %% — the query is %-formatted
        data = {r["account_id"]: r for r in _q(
            "SELECT jl.account_id, "
            f"ISNULL(SUM(CASE WHEN je.entry_date < ? OR {opn} THEN jl.debit - jl.credit "
            "ELSE 0 END),0) AS op, "
            f"ISNULL(SUM(CASE WHEN je.entry_date >= ? AND NOT {opn} THEN jl.debit "
            "ELSE 0 END),0) AS inn, "
            f"ISNULL(SUM(CASE WHEN je.entry_date >= ? AND NOT {opn} THEN jl.credit "
            "ELSE 0 END),0) AS outt "
            "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
            "WHERE je.is_posted=1 AND je.entry_date <= ? AND jl.account_id IN (%s) "
            "GROUP BY jl.account_id" % ",".join("?" * len(ids)),
            (date_from, date_from, date_from, date_to, *ids))}
        for a in accs:
            d = data.get(a["account_id"], {})
            op, i, o = _f(d.get("op")), _f(d.get("inn")), _f(d.get("outt"))
            if not any(abs(x) > 0.004 for x in (op, i, o)):
                continue
            rows.append({"code": a["account_code"], "name": a["account_name_en"],
                         "opening": round(op, 2), "in": round(i, 2), "out": round(o, 2),
                         "closing": round(op + i - o, 2)})
    cols = [col("code", "Code", width=80), col("name", "Account", width=240),
            col("opening", "Opening", "money"), col("in", "Money in", "money"),
            col("out", "Money out", "money"), col("closing", "Closing", "money")]
    tot = _sum_totals(rows, ("opening", "in", "out", "closing"), "name", "TOTAL")
    cash = sum(r["closing"] for r in rows if r["code"] == _default_code("cash"))
    return _result("Cash & Bank Balances", cols, rows, tot,
                   [("Opening", tot["opening"], "money"), ("Money in", tot["in"], "money"),
                    ("Money out", tot["out"], "money"), ("Closing", tot["closing"], "money"),
                    ("Cash in hand", cash, "money")],
                   ["Money moved between your own cash / bank accounts shows as in and out. "
                    "Opening balances (Accounts ▸ Opening Balances) are in the Opening column."])


def _expense_ids():
    """Expense accounts that are NOT cost of sales (they are in Gross Profit)."""
    try:
        from modules import account_defaults as ad
        cos = {ad.cogs(), ad.purchase_return(), ad.stock_adjust()}
    except Exception:
        cos = set()
    return {r["account_id"]: r for r in _q(
        "SELECT account_id, account_code, account_name_en FROM accounts WHERE account_type='expense'")
            if r["account_id"] not in cos and not str(r["account_code"]).startswith("51")}


@report("day_report", GROUP_ACC, "Day Report",
        desc="Each day: sales, returns, purchases, expenses, money received / paid, "
             "cash in / out and the closing cash & bank balance.")
def r_day_report(date_from, date_to, **_):
    days = defaultdict(lambda: defaultdict(float))
    for h in _sales_headers(date_from, date_to):
        d = _d(h["invoice_date"])
        if _is_credit_note(h):
            days[d]["returns"] += _f(h.get("total_amount"))
        else:
            days[d]["sales"] += _f(h.get("total_amount"))
            days[d]["bills"] += 1
    for h in _return_headers("sales", date_from, date_to):
        days[_d(h["return_date"])]["returns"] += _f(h.get("total_amount"))
    for h in _purchase_headers(date_from, date_to):
        days[_d(h["invoice_date"])]["purchases"] += _f(h.get("total_amount"))
    for h in _return_headers("purchase", date_from, date_to):
        days[_d(h["return_date"])]["purchases"] -= _f(h.get("total_amount"))
    for p in _q("SELECT payment_date, payment_type, amount FROM payments "
                "WHERE payment_date BETWEEN ? AND ?", (date_from, date_to)):
        k = "received" if str(p.get("payment_type") or "").lower() == "receipt" else "paid"
        days[_d(p["payment_date"])][k] += _f(p.get("amount"))
    exp = _expense_ids()
    if exp:
        for r in _q("SELECT je.entry_date, ISNULL(SUM(jl.debit - jl.credit),0) AS a "
                    "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
                    "WHERE je.is_posted=1 AND je.entry_date BETWEEN ? AND ? AND jl.account_id IN (%s) "
                    "GROUP BY je.entry_date" % ",".join("?" * len(exp)),
                    (date_from, date_to, *exp.keys())):
            days[_d(r["entry_date"])]["expenses"] += _f(r["a"])
    cash_ids = [a["account_id"] for a in _cash_accounts()]
    opening = 0.0
    if cash_ids:
        ph = ",".join("?" * len(cash_ids))
        opening = _f((_q1("SELECT ISNULL(SUM(jl.debit - jl.credit),0) AS b FROM journal_lines jl "
                          "JOIN journal_entries je ON je.journal_id=jl.journal_id WHERE je.is_posted=1 "
                          "AND je.entry_date < ? AND jl.account_id IN (%s)" % ph,
                          (date_from, *cash_ids)) or {}).get("b"))
        for r in _q("SELECT je.entry_date, je.journal_id, je.source_type, "
                    "ISNULL(SUM(jl.debit - jl.credit),0) AS n "
                    "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
                    "WHERE je.is_posted=1 AND je.entry_date BETWEEN ? AND ? AND jl.account_id IN (%s) "
                    "GROUP BY je.entry_date, je.journal_id, je.source_type" % ph,
                    (date_from, date_to, *cash_ids)):
            n = _f(r["n"])
            if str(r.get("source_type") or "").startswith("opening"):
                opening += n          # the opening cash / bank is the start, not money in
                continue
            dd = days[_d(r["entry_date"])]
            if n > 0:
                dd["cash_in"] += n
            elif n < 0:
                dd["cash_out"] -= n
    rows = []
    bal = opening
    for d in sorted(k for k in days if k):
        v = days[d]
        bal += v["cash_in"] - v["cash_out"]
        rows.append({"date": d, "bills": int(v["bills"]), "sales": round(v["sales"], 2),
                     "returns": round(v["returns"], 2),
                     "net": round(v["sales"] - v["returns"], 2),
                     "purchases": round(v["purchases"], 2), "expenses": round(v["expenses"], 2),
                     "received": round(v["received"], 2), "paid": round(v["paid"], 2),
                     "cash_in": round(v["cash_in"], 2), "cash_out": round(v["cash_out"], 2),
                     "balance": round(bal, 2)})
    cols = [col("date", "Date", width=90), col("bills", "Bills", "int", 55),
            col("sales", "Sales", "money", 95), col("returns", "Returns", "money", 85),
            col("net", "Net sales", "money", 95), col("purchases", "Purchases", "money", 95),
            col("expenses", "Expenses", "money", 90), col("received", "Received", "money", 90),
            col("paid", "Paid out", "money", 90), col("cash_in", "Cash in", "money", 90),
            col("cash_out", "Cash out", "money", 90), col("balance", "Cash & bank", "money", 100)]
    tot = _sum_totals(rows, ("bills", "sales", "returns", "net", "purchases", "expenses",
                             "received", "paid", "cash_in", "cash_out"), "date", "TOTAL")
    tot["balance"] = round(bal, 2)
    return _result("Day Report", cols, rows, tot,
                   [("Net sales", tot["net"], "money"), ("Purchases", tot["purchases"], "money"),
                    ("Expenses", tot["expenses"], "money"),
                    ("Cash & bank at start", opening, "money"),
                    ("Cash & bank at end", bal, "money")],
                   ["Sales and purchases include VAT. Purchases are after purchase returns. "
                    "Expenses come from the books (rent, salaries … not cost of sales). "
                    "Cash in / out: cash, bank and card accounts; moving money between them "
                    "is not counted. Opening cash / bank (Accounts ▸ Opening Balances) is "
                    "part of the start balance."])


def _aging_report(kind, date_to):
    from modules import aging as AG
    rows = AG.party_aging(kind, date_to)
    who = "Customer" if kind == "customer" else "Supplier"
    out = [{"code": r["code"], "name": r["name"], "phone": r["phone"],
            "terms": r["terms"] or None, "not_due": r["not_due"], "b30": r["b30"],
            "b60": r["b60"], "b90": r["b90"], "b90p": r["b90p"],
            "advance": r["advance"] or None, "balance": r["balance"],
            "oldest": r["oldest"] or None} for r in rows]
    cols = [col("code", "Code", width=80), col("name", who, width=200),
            col("phone", "Phone", width=100), col("terms", "Terms", "int", 50),
            col("not_due", "Not due", "money", 85), col("b30", "1-30", "money", 80),
            col("b60", "31-60", "money", 80), col("b90", "61-90", "money", 80),
            col("b90p", "90+", "money", 80), col("advance", "Advance", "money", 80),
            col("balance", "Balance", "money", 95), col("oldest", "Oldest (days)", "int", 70)]
    t = AG.totals(rows)
    tot = {k: t.get(k, 0.0) for k in ("not_due", "b30", "b60", "b90", "b90p", "advance", "balance")}
    tot["name"] = "TOTAL"
    over = t["b30"] + t["b60"] + t["b90"] + t["b90p"]
    title = "Receivable Aging" if kind == "customer" else "Payable Aging"
    return _result(title, cols, out, tot,
                   [(("Customers owe you" if kind == "customer" else "You owe suppliers"),
                     t["balance"] + t["advance"], "money"),
                    ("Not yet due", t["not_due"], "money"), ("Overdue", over, "money"),
                    ("Over 90 days", t["b90p"], "money"), ("Advances", t["advance"], "money")],
                   ["Balances from each party's ledger. Age = days past the due date "
                    "(bill date + payment terms). A payment settled against bills closes "
                    "those bills; money not tied to a bill settles the oldest bills first."])


@report("receivable_aging", GROUP_ACC, "Receivable Aging", params=("as_of",),
        desc="What customers owe, by how many days it is overdue.")
def r_receivable_aging(date_to, **_):
    return _aging_report("customer", date_to)


@report("payable_aging", GROUP_ACC, "Payable Aging", params=("as_of",),
        desc="What you owe suppliers, by how many days it is overdue.")
def r_payable_aging(date_to, **_):
    return _aging_report("supplier", date_to)


def _bills_report(kind, date_to, party=""):
    from modules import aging as AG
    bills = [b for b in AG.open_bills(kind, date_to)
             if _match(party, b["party"], b["code"])]
    rows = [{"date": b["date"], "due": b["due"], "party": b["party"], "type": b["type"],
             "ref": b["ref"], "amount": b["amount"], "open": b["open"],
             "days": max(0, b["overdue"]) or None} for b in bills]
    who = "Customer" if kind == "customer" else "Supplier"
    cols = [col("date", "Date", width=90), col("due", "Due", width=90),
            col("party", who, width=200), col("type", "Document", width=120),
            col("ref", "Ref", width=110), col("amount", "Bill amount", "money", 95),
            col("open", "Still open", "money", 95), col("days", "Days overdue", "int", 80)]
    tot = _sum_totals(rows, ("amount", "open"), "party", "TOTAL")
    over = sum(r["open"] for r in rows if r["days"])
    title = ("Customer Outstanding Bills" if kind == "customer" else "Supplier Outstanding Bills")
    return _result(title, cols, rows, tot,
                   [("Open bills", len(rows), "int"), ("Still open", tot["open"], "money"),
                    ("Overdue", over, "money")],
                   ["A receipt / payment settled against bills (payment screen) and a "
                    "return made from a bill close those bills; money not tied to a bill "
                    "settles the oldest bills first."])


@report("receivable_bills", GROUP_ACC, "Customer Outstanding Bills", params=("as_of", "party"),
        desc="Every customer bill that is not fully paid, with days overdue.")
def r_receivable_bills(date_to, party="", **_):
    return _bills_report("customer", date_to, party)


@report("payable_bills", GROUP_ACC, "Supplier Outstanding Bills", params=("as_of", "party"),
        desc="Every supplier bill you have not fully paid, with days overdue.")
def r_payable_bills(date_to, party="", **_):
    return _bills_report("supplier", date_to, party)


@report("expense_summary", GROUP_ACC, "Expense Summary",
        desc="Every expense in the books per account (and per month): rent, salaries, "
             "bills … — not cost of sales.")
def r_expense_summary(date_from, date_to, **_):
    exp = _expense_ids()
    data = []
    if exp:
        data = _q("SELECT jl.account_id, je.entry_date, je.source_type, "
                  "ISNULL(jl.debit,0) - ISNULL(jl.credit,0) AS a "
                  "FROM journal_lines jl JOIN journal_entries je ON je.journal_id=jl.journal_id "
                  "WHERE je.is_posted=1 AND je.entry_date BETWEEN ? AND ? "
                  "AND jl.account_id IN (%s)" % ",".join("?" * len(exp)),
                  (date_from, date_to, *exp.keys()))
    months = sorted({str(_d(r["entry_date"]))[:7] for r in data})
    by_month = len(months) <= 12
    acc = defaultdict(lambda: defaultdict(float))
    count = defaultdict(int)
    for r in data:
        a = r["account_id"]
        acc[a]["total"] += _f(r["a"])
        count[a] += 1
        if by_month:
            acc[a][str(_d(r["entry_date"]))[:7]] += _f(r["a"])
    grand = sum(v["total"] for v in acc.values())
    rows = []
    for a, v in sorted(acc.items(), key=lambda x: -x[1]["total"]):
        if abs(v["total"]) < 0.005:
            continue
        row = {"code": exp[a]["account_code"], "name": exp[a]["account_name_en"],
               "count": count[a], "total": round(v["total"], 2),
               "pct": round(100.0 * v["total"] / grand, 1) if grand else 0.0}
        for m in (months if by_month else []):
            row["m" + m] = round(v.get(m, 0.0), 2) or None
        rows.append(row)
    cols = [col("code", "Code", width=70), col("name", "Expense account", width=200),
            col("count", "Entries", "int", 55)]
    if by_month and len(months) > 1:
        cols += [col("m" + m, _dt.date(int(m[:4]), int(m[5:]), 1).strftime("%b %Y"), "money", 85)
                 for m in months]
    cols += [col("total", "Total", "money", 95), col("pct", "% of all", "pct", 65)]
    keys = ["count", "total"] + (["m" + m for m in months] if by_month and len(months) > 1 else [])
    tot = _sum_totals(rows, keys, "name", "TOTAL")
    tot["pct"] = 100.0 if rows else 0.0
    top = rows[0]["name"] if rows else "—"
    return _result("Expense Summary", cols, rows, tot,
                   [("Total expenses", tot["total"], "money"), ("Accounts", len(rows), "int"),
                    ("Biggest", top, "text")],
                   ["From the books: expense vouchers, journal vouchers, salaries, "
                    "depreciation, rounding … Cost of sales is in Profit & Loss."])


@report("opening_balances", GROUP_ACC, "Opening Balances", params=("as_now",),
        desc="Everything entered as opening: accounts, stock, customers and suppliers.")
def r_opening_balances(**_):
    from modules import opening_balances as OB
    rows = []

    def head(t):
        rows.append({"code": "", "name": t, "qty": None, "dr": None, "cr": None, "_style": "head"})
    acc = [r for r in OB.account_rows() if r["debit"] or r["credit"]]
    if acc:
        head("Accounts")
        for r in acc:
            rows.append({"code": r["code"], "name": "    " + r["name"], "qty": None,
                         "dr": r["debit"] or None, "cr": r["credit"] or None})
    stock = [r for r in OB.stock_rows() if abs(r["opening"]) > 1e-9]
    if stock:
        head("Opening stock")
        for r in stock:
            rows.append({"code": r["code"], "name": f"    {r['name']}  @ {r['cost']:,.2f}",
                         "qty": r["opening"], "dr": r["value"] or None, "cr": None})
    for kind, title, side in (("customer", "Customers (owe you)", "dr"),
                              ("supplier", "Suppliers (you owe)", "cr")):
        ps = [r for r in OB.party_rows(kind) if abs(r["opening"]) >= 0.005]
        if ps:
            head(title)
            for r in ps:
                a = r["opening"]
                d = {"code": r["code"], "name": "    " + r["name"], "qty": None,
                     "dr": None, "cr": None}
                if (side == "dr") == (a > 0):
                    d["dr"] = abs(a)
                else:
                    d["cr"] = abs(a)
                rows.append(d)
    s = OB.summary()
    tdr = sum(_f(r["dr"]) for r in rows if not r.get("_style"))
    tcr = sum(_f(r["cr"]) for r in rows if not r.get("_style"))
    rows.append({"code": "3400", "name": "Opening Balance Equity (the difference)",
                 "qty": None, "dr": round(tcr - tdr, 2) if tcr > tdr else None,
                 "cr": round(tdr - tcr, 2) if tdr > tcr else None, "_style": "sub"})
    # moved into Owner's Capital since (Opening Balances ▸ Move to Owner's Capital)
    moved = 0.0
    try:
        from modules import account_defaults as ad
        moved = round(_f(_q1(
            "SELECT SUM(jl.debit) - SUM(jl.credit) AS m FROM journal_lines jl "
            "JOIN journal_entries je ON je.journal_id=jl.journal_id "
            "WHERE je.reference='OPEN-CAP' AND je.is_posted=1 AND jl.account_id=?",
            (ad.opening_equity(),)).get("m")), 2)
    except Exception:
        moved = 0.0
    if abs(moved) >= 0.005:
        rows.append({"code": "3100", "qty": None, "dr": None, "cr": None, "_style": "sub",
                     "name": f"    → {abs(moved):,.2f} since moved to Owner's Capital"})
    cols = [col("code", "Code", width=90), col("name", "Particulars", width=300),
            col("qty", "Qty", "qty", 70), col("dr", "Debit", "money", 110),
            col("cr", "Credit", "money", 110)]
    tot = {"name": "TOTAL", "dr": round(max(tdr, tcr), 2), "cr": round(max(tdr, tcr), 2)}
    return _result("Opening Balances", cols, rows, tot,
                   [("Opening date", s["date"], "text"),
                    ("Stock value", s["stock_value"], "money"),
                    ("Customers owe", s["receivable"], "money"),
                    ("You owe suppliers", s["payable"], "money"),
                    ("Opening Balance Equity", s["equity"], "money"),
                    ("Left in Opening Balance Equity", OB.equity_balance(), "money")],
                   ["Enter or change them in Accounts ▸ Opening Balances. "
                    "Move to Owner's Capital closes Opening Balance Equity."],
                   subtitle=f"Opening date {s['date']}")


@report("stock_transactions", GROUP_INV, "Stock Transactions",
        params=("period", "product", "category"),
        desc="Every stock in / out of every product: purchases, sales, returns, "
             "counts, adjustments, opening stock.")
def r_stock_transactions(date_from, date_to, product="", category="", **_):
    cats = _categories()
    prods = _products()
    rows = []
    for m in sorted(_movements(date_from, date_to), key=lambda x: (x["date"], x["type"])):
        p = prods.get(m["pid"]) or {}
        if not _match(product, p.get("product_code"), p.get("product_name_en"),
                      p.get("product_name_ar"), p.get("barcode")):
            continue
        cat = cats.get(p.get("category_id"), "") or "(no category)"
        if category and category != "All" and cat != category:
            continue
        c = _cost(p)
        rows.append({"date": m["date"], "type": m["type"], "ref": m["ref"],
                     "code": p.get("product_code") or "", "name": p.get("product_name_en") or "",
                     "party": m["party"], "qin": round(m["qin"], 3) or None,
                     "qout": round(m["qout"], 3) or None,
                     "value": round((m["qin"] - m["qout"]) * c, 2)})
    cols = [col("date", "Date", width=85), col("type", "Type", width=110),
            col("ref", "Ref", width=100), col("code", "Code", width=70),
            col("name", "Product", width=190), col("party", "Party / note", width=170),
            col("qin", "In", "qty", 65), col("qout", "Out", "qty", 65),
            col("value", "Value (cost)", "money", 90)]
    tot = _sum_totals(rows, ("qin", "qout", "value"), "name", "TOTAL")
    return _result("Stock Transactions", cols, rows, tot,
                   [("Movements", len(rows), "int"), ("In", tot["qin"], "qty"),
                    ("Out", tot["qout"], "qty"), ("Net value", tot["value"], "money")],
                   ["Quantities in base units, valued at today's average cost."])


# ── plain-text / CSV helpers (connector, tests) ───────────────
def as_table(res, max_rows=None):
    cols = res["columns"]
    out = [[c["label"] for c in cols]]
    rows = res["rows"][:max_rows] if max_rows else res["rows"]
    for r in rows:
        out.append([r.get(c["key"]) for c in cols])
    if res.get("totals"):
        out.append([res["totals"].get(c["key"]) for c in cols])
    return out


# ════════════════════════════════════════════════════════════
#  v6.0 — shifts, stock per branch, batches & expiry
# ════════════════════════════════════════════════════════════
@report("cashier_watch", GROUP_SALES, "Cashier Watch",
        desc="Per cashier: bills, sales, discounts, returns, items removed before payment, "
             "bills cleared, documents deleted, cash short / over — with a warning when one "
             "cashier is far above the others.")
def r_cashier_watch(date_from, date_to, **_):
    from modules import cashier_watch as CW
    rows = []
    for x in CW.report(date_from, date_to):
        rows.append({"cashier": x["cashier"], "bills": x["bills"], "sales": x["sales"],
                     "discount": x["discount"], "discount_pct": x["discount_pct"],
                     "returns": x["returns"], "returns_amt": x["returns_amt"],
                     "voids": x["voids"], "clears": x["clears"], "voids_amt": x["voids_amt"],
                     "deletes": x["deletes"], "cash_short": x["cash_short"] or None,
                     "cash_over": x["cash_over"] or None,
                     "flags": "  ·  ".join(x["flags"]) or "✔",
                     "_style": "hi" if x["flags"] else None})
    cols = [col("cashier", "Cashier", width=130), col("flags", "Warnings", width=340),
            col("bills", "Bills", "int", 60),
            col("sales", "Sales", "money"), col("discount", "Discounts", "money", 95),
            col("discount_pct", "Disc. %", "pct", 65), col("returns", "Returns", "int", 65),
            col("returns_amt", "Returned", "money", 95), col("voids", "Items removed", "int", 95),
            col("clears", "Bills cleared", "int", 90), col("voids_amt", "Removed SAR", "money", 95),
            col("deletes", "Deleted docs", "int", 90), col("cash_short", "Cash short", "money", 90),
            col("cash_over", "Cash over", "money", 85)]
    flagged = sum(1 for r in rows if r["flags"] != "✔")
    return _result("Cashier Watch", cols, rows, None,
                   [("Cashiers", len(rows), "int"), ("With warnings", flagged, "int"),
                    ("Items removed", sum(r["voids"] for r in rows), "int"),
                    ("Discounts", sum(r["discount"] for r in rows), "money")],
                   ["A warning = far above the other cashiers (about 2×) or above a fixed limit; "
                    "cash short above %.0f SAR (Settings ▸ POS). Check it — it is not proof of "
                    "anything." % CW.short_limit()])


@report("cashier_shifts", GROUP_SALES, "Cashier Shifts",
        desc="Every cashier shift: opening cash, bills, sales by tender, refunds, expected cash, "
             "counted cash and the difference (short / over).")
def r_cashier_shifts(date_from, date_to, **_):
    from modules import shifts as SH
    rows = []
    for s in SH.history(date_from, date_to, 1000):
        sm = SH.summary(s)
        ten = sm.get("tenders") or {}
        rows.append({"id": s["shift_id"], "opened": str(s.get("opened_at") or "")[:16],
                     "closed": str(s.get("closed_at") or "")[:16] or "open",
                     "cashier": s.get("username") or "", "float": sm.get("opening_float"),
                     "bills": sm.get("bills"), "sales": sm.get("sales"),
                     "cash": ten.get("cash", 0.0),
                     "other": round(sum(v for k, v in ten.items() if k != "cash"), 2),
                     "refunds": sm.get("cash_refunds"), "expected": sm.get("expected_cash"),
                     "counted": _f(s.get("counted_cash")) if s.get("counted_cash") is not None else None,
                     "diff": _f(s.get("difference")) if s.get("difference") is not None else None})
    cols = [col("id", "Shift", width=55), col("opened", "Opened", width=120),
            col("closed", "Closed", width=120), col("cashier", "Cashier", width=100),
            col("float", "Opening cash", "money", 95), col("bills", "Bills", "int", 55),
            col("sales", "Sales", "money", 95), col("cash", "Cash", "money", 90),
            col("other", "Card / other", "money", 95), col("refunds", "Cash refunds", "money", 90),
            col("expected", "Expected cash", "money", 100), col("counted", "Counted", "money", 90),
            col("diff", "Short (−) / over (+)", "money", 110)]
    tot = _sum_totals(rows, ("bills", "sales", "cash", "other", "refunds", "diff"), "cashier", "TOTAL")
    short = sum(1 for r in rows if (r["diff"] or 0) < -0.004)
    return _result("Cashier Shifts", cols, rows, tot,
                   [("Shifts", len(rows), "int"), ("Short shifts", short, "int"),
                    ("Total difference", tot.get("diff") or 0.0, "money")])


@report("stock_by_branch", GROUP_INV, "Stock by Branch", params=("as_now", "category"),
        desc="Stock of every product in each branch (transfers between branches included).")
def r_stock_by_branch(category="", **_):
    from modules import stock_branch as SB
    brs = SB.branches()
    cats = _categories()
    items = [p for p in _stock_items()
             if not category or category == "All" or (cats.get(p.get("category_id"), "") or "(no category)") == category]
    by = SB.qty_by_branch([p["product_id"] for p in items]) if items else {}
    rows = []
    for p in items:
        r = {"code": p.get("product_code") or "", "name": p.get("product_name_en") or "",
             "total": round(_f(p.get("stock_quantity")), 3)}
        for b in brs:
            r["b%s" % b["branch_id"]] = round(_f(by.get(p["product_id"], {}).get(b["branch_id"])), 3)
        rows.append(r)
    rows.sort(key=lambda r: r["name"])
    cols = [col("code", "Code", width=80), col("name", "Product", width=220),
            col("total", "Total", "qty", 80)]
    cols += [col("b%s" % b["branch_id"], (b.get("name_en") or str(b["branch_id"]))[:18], "qty", 95)
             for b in brs]
    tot = _sum_totals(rows, ["total"] + ["b%s" % b["branch_id"] for b in brs], "name", "TOTAL")
    notes = [] if len(brs) > 1 else ["Only one branch — every product's stock is in it."]
    return _result("Stock by Branch", cols, rows, tot, [("Branches", len(brs), "int"),
                                                         ("Products", len(rows), "int")],
                   notes, subtitle="Current stock")


@report("expiry_report", GROUP_INV, "Batches & Expiry", params=("as_now",),
        desc="Batches still in stock with their expiry dates — expired and soon-to-expire first, "
             "with their value at cost.")
def r_expiry(**_):
    from modules import batches as BT
    rows = []
    for b in BT.list_batches():
        d = b.get("days_left")
        rows.append({"code": b.get("product_code") or "", "name": b.get("product_name_en") or "",
                     "batch": b.get("batch_no") or "", "expiry": str(b.get("expiry_date") or "")[:10],
                     "days": d, "qty": round(_f(b.get("qty_base")), 3),
                     "value": round(_f(b.get("qty_base")) * _f(b.get("unit_cost")), 2),
                     "state": "EXPIRED" if (d is not None and d < 0) else
                              ("soon" if (d is not None and d <= 30) else "")})
    cols = [col("code", "Code", width=80), col("name", "Product", width=220),
            col("batch", "Batch", width=90), col("expiry", "Expiry", width=90),
            col("days", "Days left", "int", 70), col("qty", "Quantity", "qty", 80),
            col("value", "Value (cost)", "money", 100), col("state", "", width=70)]
    tot = _sum_totals(rows, ("qty", "value"), "name", "TOTAL")
    gone = [r for r in rows if r["state"] == "EXPIRED"]
    return _result("Batches & Expiry", cols, rows, tot,
                   [("Batches", len(rows), "int"), ("Expired", len(gone), "int"),
                    ("Expired value", round(sum(r["value"] for r in gone), 2), "money")],
                   subtitle="Current stock")


# ════════════════════════════════════════════════════════════
#  EXPORT  (Excel / PDF) — used by the Reports screen
# ════════════════════════════════════════════════════════════
def fmt(v, kind):
    if v is None or v == "":
        return ""
    if kind == "money":
        try:
            x = float(v)
            return f"({abs(x):,.2f})" if x < 0 else f"{x:,.2f}"
        except Exception:
            return str(v)
    if kind == "qty":
        try:
            x = float(v)
            return f"{x:,.3f}".rstrip("0").rstrip(".") if x % 1 else f"{x:,.0f}"
        except Exception:
            return str(v)
    if kind == "pct":
        try:
            return f"{float(v):.1f}%"
        except Exception:
            return str(v)
    if kind == "int":
        try:
            return f"{int(round(float(v))):,}"
        except Exception:
            return str(v)
    return str(v)


def _company():
    try:
        return _q1("SELECT * FROM company WHERE company_id=1") or _q1("SELECT * FROM company") or {}
    except Exception:
        return {}


def to_excel(res, path, company=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    company = company if company is not None else _company()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (res.get("title") or "Report")[:31].replace("/", "-")
    cols = res["columns"]
    n = max(1, len(cols))
    r = 1
    ws.cell(r, 1, company.get("company_name_en") or "DayBill ERP").font = Font(bold=True, size=14, color="1F2A44")
    r += 1
    ws.cell(r, 1, res.get("title", "")).font = Font(bold=True, size=12)
    r += 1
    ws.cell(r, 1, f"{res.get('subtitle', '')}    Generated {_dt.datetime.now():%Y-%m-%d %H:%M}").font = Font(size=9, color="6B7280")
    r += 2
    for lbl, val, kind in res.get("summary") or []:
        ws.cell(r, 1, lbl).font = Font(bold=True)
        c = ws.cell(r, 2, round(float(val), 4) if kind in NUMERIC and isinstance(val, (int, float))
                    else val)
        if kind == "money":
            c.number_format = '#,##0.00;(#,##0.00)'
        r += 1
    if res.get("summary"):
        r += 1
    head_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="D9D9D9")
    for i, cdef in enumerate(cols, 1):
        c = ws.cell(r, i, cdef["label"])
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(r + 1, 1)
    r += 1
    fills = {"head": PatternFill("solid", fgColor="E8EEF7"), "total": PatternFill("solid", fgColor="F2F2F2"),
             "grand": PatternFill("solid", fgColor="DDEBF7"), "hi": PatternFill("solid", fgColor="F3F8FD"),
             "sub": PatternFill("solid", fgColor="FAFAFA")}
    rows = list(res["rows"])
    if res.get("totals"):
        t = dict(res["totals"]); t["_style"] = "total"
        rows.append(t)
    for row in rows:
        st = row.get("_style") or ""
        for i, cdef in enumerate(cols, 1):
            v = row.get(cdef["key"])
            if cdef["kind"] in NUMERIC and isinstance(v, (int, float)):
                c = ws.cell(r, i, round(float(v), 4))
                c.number_format = {"money": '#,##0.00;(#,##0.00)', "qty": '#,##0.###',
                                   "pct": '0.0"%"', "int": '#,##0'}[cdef["kind"]]
                c.alignment = Alignment(horizontal="right")
            else:
                c = ws.cell(r, i, "" if v is None else str(v))
            c.border = Border(bottom=thin)
            if st in ("head", "total", "grand"):
                c.font = Font(bold=True)
            if st in fills:
                c.fill = fills[st]
        r += 1
    for note in res.get("notes") or []:
        r += 1
        ws.cell(r, 1, note).font = Font(italic=True, size=9, color="6B7280")
    for i, cdef in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(48, cdef["width"] // 7))
    ws.page_setup.orientation = "landscape" if n > 7 else "portrait"
    ws.page_setup.fitToWidth = 1
    wb.save(path)
    return path


def _has_arabic(s):
    return any("؀" <= ch <= "ۿ" for ch in str(s))


def to_pdf(res, path, company=None):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    company = company if company is not None else _company()
    font, font_b = "Helvetica", "Helvetica-Bold"
    shape = lambda s: s  # noqa: E731
    try:
        from modules import pdf_invoice as _pi
        _pi._register_fonts()
        if getattr(_pi, "_AR_FONT", None):
            font = _pi._AR_FONT
            font_b = getattr(_pi, "_AR_FONT_BOLD", None) or font
            shape = _pi._shape_arabic
    except Exception:
        pass

    def txt(s):
        s = "" if s is None else str(s)
        return shape(s) if _has_arabic(s) else s

    cols = res["columns"]
    wide = sum(c["width"] for c in cols) > 820 or len(cols) > 7
    page = landscape(A4) if wide else A4
    doc = SimpleDocTemplate(path, pagesize=page, topMargin=14 * mm, bottomMargin=14 * mm,
                            leftMargin=10 * mm, rightMargin=10 * mm,
                            title=res.get("title", "Report"))
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontName=font_b, fontSize=15,
                        textColor=colors.HexColor("#1F2A44"), alignment=0, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName=font_b, fontSize=12, spaceAfter=2)
    sm = ParagraphStyle("sm", parent=ss["Normal"], fontName=font, fontSize=8,
                        textColor=colors.HexColor("#6B7280"))
    el = [Paragraph(txt(company.get("company_name_en") or "DayBill ERP"), h1)]
    if company.get("vat_registration_no"):
        el.append(Paragraph("VAT No. " + str(company["vat_registration_no"]), sm))
    el.append(Paragraph(txt(res.get("title", "")), h2))
    el.append(Paragraph(f"{res.get('subtitle', '')} &nbsp;&nbsp; Generated "
                        f"{_dt.datetime.now():%Y-%m-%d %H:%M}", sm))
    el.append(Spacer(1, 6))
    if res.get("summary"):
        cells = [[txt(l), fmt(v, k) if k in NUMERIC else txt(v)] for l, v, k in res["summary"]]
        pairs = []
        for i in range(0, len(cells), 3):
            row = []
            for c in cells[i:i + 3]:
                row += c
            row += [""] * (6 - len(row))
            pairs.append(row)
        st = Table(pairs, hAlign="LEFT")
        st.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                                ("FONTNAME", (1, 0), (1, -1), font_b), ("FONTNAME", (3, 0), (3, -1), font_b),
                                ("FONTNAME", (5, 0), (5, -1), font_b),
                                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#1F2A44")),
                                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3F6FB")),
                                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D3E3")),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 10)]))
        el += [st, Spacer(1, 8)]
    # column widths first: a heading or a long text that does not fit its
    # column wraps onto a second line instead of running into the next one
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from xml.sax.saxutils import escape as _esc
    avail = page[0] - 20 * mm
    tw = sum(c["width"] for c in cols) or 1
    widths = [avail * c["width"] / tw for c in cols]
    fsize = 7.8 if wide else 8.5
    hd_l = ParagraphStyle("rc_hd", parent=ss["Normal"], fontName=font_b, fontSize=fsize,
                          leading=fsize + 1.5, textColor=colors.white)
    hd_r = ParagraphStyle("rc_hdr", parent=hd_l, alignment=2)
    cl = ParagraphStyle("rc_cell", parent=ss["Normal"], fontName=font, fontSize=fsize,
                        leading=fsize + 1.5)
    cl_b = ParagraphStyle("rc_cellb", parent=cl, fontName=font_b)

    def fits(s_, w, f):
        try:
            return stringWidth(s_, f, fsize) <= w - 6
        except Exception:
            return True
    data = [[c["label"] if fits(c["label"], widths[ci], font_b) else
             Paragraph(_esc(c["label"]), hd_r if c["kind"] in NUMERIC else hd_l)
             for ci, c in enumerate(cols)]]
    styles = []
    rows = list(res["rows"])
    if res.get("totals"):
        t = dict(res["totals"]); t["_style"] = "total"
        rows.append(t)
    for i, row in enumerate(rows, 1):
        bold = row.get("_style") in ("head", "total", "grand")
        line = []
        for ci, c in enumerate(cols):
            if c["kind"] in NUMERIC:
                line.append(fmt(row.get(c["key"]), c["kind"]))
                continue
            v = txt(row.get(c["key"]))
            if v and not fits(v, widths[ci], font_b if bold else font):
                v = Paragraph(_esc(v), cl_b if bold else cl)
            line.append(v)
        data.append(line)
        st = row.get("_style")
        if st in ("head", "total", "grand"):
            styles.append(("FONTNAME", (0, i), (-1, i), font_b))
        if st == "head":
            styles.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#E8EEF7")))
        elif st == "total":
            styles.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F2F2F2")))
            styles.append(("LINEABOVE", (0, i), (-1, i), 0.6, colors.HexColor("#8497B0")))
        elif st == "grand":
            styles.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#DDEBF7")))
            styles.append(("LINEABOVE", (0, i), (-1, i), 0.8, colors.HexColor("#1F4E79")))
    tbl = Table(data, colWidths=widths, repeatRows=1)
    base = [("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), fsize),
            ("FONTNAME", (0, 0), (-1, 0), font_b),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#D9D9D9")),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for ci, c in enumerate(cols):
        if c["kind"] in NUMERIC:
            base.append(("ALIGN", (ci, 0), (ci, -1), "RIGHT"))
    tbl.setStyle(TableStyle(base + styles))
    el.append(tbl)
    for note in res.get("notes") or []:
        el += [Spacer(1, 4), Paragraph(txt(note), sm)]

    def _footer(canv, d):
        canv.saveState()
        canv.setFont("Helvetica", 7.5)
        canv.setFillColor(colors.HexColor("#6B7280"))
        canv.drawRightString(page[0] - 10 * mm, 8 * mm, f"Page {d.page}")
        canv.drawString(10 * mm, 8 * mm, "DayBill ERP · " + (res.get("title") or ""))
        canv.restoreState()
    doc.build(el, onFirstPage=_footer, onLaterPages=_footer)
    return path


# ════════════════════════════════════════════════════════════
#  v6.6 — HR & PAYROLL REPORTS
# ════════════════════════════════════════════════════════════
GROUP_HR = "HR & Payroll"


def _runs_in(f, t):
    """Payroll months whose first day falls between the first of `f`'s month and `t`."""
    f0 = f[:7] + "-01"
    out = []
    for r in _q("SELECT * FROM payroll_runs ORDER BY run_year, run_month"):
        d = "%04d-%02d-01" % (int(r["run_year"]), int(r["run_month"]))
        if f0 <= d <= t:
            out.append(r)
    return out


@report("payroll_register", GROUP_HR, "Payroll Register",
        desc="Every payslip of the months in the period: basic, allowances, overtime, bonus, gross, GOSI, "
             "absent / sick days, loan, other deductions, net.")
def r_payroll_register(date_from, date_to, **_):
    rows = []
    for run in _runs_in(date_from, date_to):
        for p in _q("SELECT * FROM payslips WHERE run_id=? ORDER BY employee_name_en", (run["run_id"],)):
            rows.append({"month": run["run_label"], "employee": p.get("employee_name_en") or "",
                         "basic": _f(p.get("basic_salary")),
                         "allow": _f(p.get("housing_allowance")) + _f(p.get("transport_allowance"))
                         + _f(p.get("other_allowance")),
                         "overtime": _f(p.get("overtime_amount")), "bonus": _f(p.get("bonus_amount")),
                         "gross": _f(p.get("gross_salary")), "gosi": _f(p.get("gosi_employee")),
                         "absent": _f(p.get("absence_deduction")) + _f(p.get("sick_deduction")),
                         "loan": _f(p.get("loan_deduction")), "other": _f(p.get("other_deduction")),
                         "net": _f(p.get("net_salary")), "status": run.get("status") or "draft"})
    cols = [col("month", "Month", width=110), col("employee", "Employee", width=170),
            col("basic", "Basic", "money", 85), col("allow", "Allowances", "money", 90),
            col("overtime", "Overtime", "money", 80), col("bonus", "Bonus", "money", 75),
            col("gross", "Gross", "money", 95), col("gosi", "GOSI", "money", 80),
            col("absent", "Absent / sick", "money", 90), col("loan", "Loan", "money", 75),
            col("other", "Other", "money", 70), col("net", "Net", "money", 95), col("status", "Status", width=70)]
    tot = _sum_totals(rows, ["basic", "allow", "overtime", "bonus", "gross", "gosi", "absent", "loan", "other",
                             "net"], "employee")
    return _result("Payroll Register", cols, rows, tot,
                   [("Payslips", len(rows), "int"), ("Gross", tot["gross"], "money"), ("Net", tot["net"], "money")])


@report("gosi_report", GROUP_HR, "GOSI Contributions", needs="country:SA",
        desc="Per employee and month: the GOSI wage (basic + housing, capped), the employee share, the employer "
             "share and the total to pay GOSI.")
def r_gosi(date_from, date_to, **_):
    from modules import payroll_engine as PE
    emps = {r["employee_id"]: r for r in _q("SELECT employee_id, employee_code, national_id, is_saudi FROM employees")}
    rows = []
    for run in _runs_in(date_from, date_to):
        for p in _q("SELECT * FROM payslips WHERE run_id=? ORDER BY employee_name_en", (run["run_id"],)):
            e = emps.get(p["employee_id"], {})
            emp_s, er_s = _f(p.get("gosi_employee")), _f(p.get("gosi_employer"))
            rows.append({"month": run["run_label"], "code": e.get("employee_code") or "",
                         "employee": p.get("employee_name_en") or "", "id": e.get("national_id") or "",
                         "saudi": "Saudi" if e.get("is_saudi") else "Non-Saudi",
                         "base": PE.gosi_contributory_wage(_f(p.get("basic_salary")), _f(p.get("housing_allowance"))),
                         "employee_share": emp_s, "employer_share": er_s, "total": round(emp_s + er_s, 2),
                         "paid": "✔ " + str(run.get("gosi_paid_at"))[:10] if run.get("gosi_paid_at") else ""})
    cols = [col("month", "Month", width=110), col("code", "Code", width=70), col("employee", "Employee", width=170),
            col("id", "Iqama / ID", width=105), col("saudi", "Nationality", width=85),
            col("base", "GOSI wage", "money", 95), col("employee_share", "Employee", "money", 90),
            col("employer_share", "Employer", "money", 90), col("total", "Total", "money", 95),
            col("paid", "Paid", width=95)]
    tot = _sum_totals(rows, ["base", "employee_share", "employer_share", "total"], "employee")
    return _result("GOSI Contributions", cols, rows, tot,
                   [("Employee share", tot["employee_share"], "money"), ("Employer share", tot["employer_share"], "money"),
                    ("To pay GOSI", tot["total"], "money")],
                   ["GOSI wage = basic + housing, up to the cap in Settings ▸ HR & Payroll. Saudis: pension + SANED "
                    "(+ hazards for the employer); others: 2 % employer (hazards)."])


@report("leave_balances", GROUP_HR, "Leave Balances", params=("as_of",),
        desc="Annual leave per employee: days a year now, earned, taken, left and what the days left are worth.")
def r_leave_balances(date_from, date_to, **_):
    from modules import hr_time as HT, payroll_engine as PE
    rows = []
    for e in _q("SELECT * FROM employees WHERE is_active=1 ORDER BY full_name_en"):
        b = HT.annual_balance(e["employee_id"], on=date_to) if HT.carry() else \
            HT.annual_balance(e["employee_id"], int(date_to[:4]))
        left = _f(b.get("balance"))
        rows.append({"code": e.get("employee_code") or "", "employee": e.get("full_name_en") or "",
                     "join": _d(e.get("join_date")), "per_year": b.get("entitlement"),
                     "earned": _f(b.get("accrued")), "taken": _f(b.get("used")), "left": left,
                     "value": round(PE.actual_wage(e, "actual") / 30.0 * max(left, 0), 2),
                     "_style": "hi" if left < 0 else None})
    cols = [col("code", "Code", width=70), col("employee", "Employee", width=180), col("join", "Joined", "date", 95),
            col("per_year", "Days a year", "int", 85), col("earned", "Earned", "qty", 80),
            col("taken", "Taken", "qty", 75), col("left", "Left", "qty", 75), col("value", "Value (SAR)", "money", 100)]
    tot = _sum_totals(rows, ["earned", "taken", "left", "value"], "employee")
    return _result("Leave Balances", cols, rows, tot,
                   [("Employees", len(rows), "int"), ("Days left", tot["left"], "qty"),
                    ("Worth", tot["value"], "money")],
                   ["Earned from the join date (or the balance typed in the employee) — 21 days a year, 30 after "
                    "5 years (art. 109). Unused days are paid at the end of service (art. 111)."
                    if HT.carry() else "Per calendar year (Settings ▸ HR: unused leave does not carry on)."])


@report("attendance_summary", GROUP_HR, "Attendance Summary",
        desc="Per employee: days present, late, absent, on leave, off, hours worked and overtime hours.")
def r_attendance(date_from, date_to, **_):
    from modules import payroll_engine as PE
    std = PE.daily_hours()
    rows = []
    names = {r["employee_id"]: r for r in _q("SELECT employee_id, employee_code, full_name_en FROM employees")}
    agg = {}
    for r in _q("SELECT employee_id, status, hours FROM attendance WHERE work_date BETWEEN ? AND ?",
                (date_from, date_to)):
        a = agg.setdefault(r["employee_id"], {"present": 0, "late": 0, "absent": 0, "leave": 0, "off": 0,
                                              "hours": 0.0, "ot": 0.0})
        st = r.get("status") or "present"
        if st in a:
            a[st] += 1
        h = r.get("hours")
        if h is not None:
            a["hours"] += _f(h)
            if st in ("present", "late"):
                a["ot"] += max(0.0, _f(h) - std)
            elif st == "off":
                a["ot"] += _f(h)
    for eid, a in agg.items():
        e = names.get(eid, {})
        rows.append({"code": e.get("employee_code") or "", "employee": e.get("full_name_en") or str(eid),
                     "present": a["present"], "late": a["late"], "absent": a["absent"], "leave": a["leave"],
                     "off": a["off"], "hours": round(a["hours"], 2), "ot": round(a["ot"], 2),
                     "_style": "hi" if a["absent"] else None})
    rows.sort(key=lambda x: x["employee"])
    cols = [col("code", "Code", width=70), col("employee", "Employee", width=180),
            col("present", "Present", "int", 70), col("late", "Late", "int", 60), col("absent", "Absent", "int", 65),
            col("leave", "On leave", "int", 70), col("off", "Day off", "int", 65),
            col("hours", "Hours", "qty", 75), col("ot", "Overtime h", "qty", 85)]
    tot = _sum_totals(rows, ["present", "late", "absent", "leave", "off", "hours", "ot"], "employee")
    return _result("Attendance Summary", cols, rows, tot,
                   [("Employees", len(rows), "int"), ("Absent days", tot["absent"], "int"),
                    ("Overtime hours", tot["ot"], "qty")],
                   ["Overtime = hours above %g a day on working days + every hour on a day off (art. 107)." % std])


@report("hr_documents", GROUP_HR, "Employee Documents Expiry", params=("as_now",),
        desc="Every Iqama / ID, passport, medical insurance and contract date of the active employees — "
             "expired first.")
def r_hr_documents(date_from, date_to, **_):
    from modules import hr_time as HT
    alert = HT.doc_alert_days()
    rows = []
    for x in HT.expiring(include_all=True):
        rows.append({"code": x["code"], "employee": x["name"], "doc": x["doc"], "date": x["date"],
                     "left": x["days_left"],
                     "state": "EXPIRED" if x["days_left"] < 0 else ("renew soon" if x["days_left"] <= alert else "ok"),
                     "_style": "hi" if x["days_left"] <= alert else None})
    cols = [col("code", "Code", width=70), col("employee", "Employee", width=190), col("doc", "Document", width=140),
            col("date", "Expires", "date", 100), col("left", "Days left", "int", 85), col("state", "", width=110)]
    due = sum(1 for r in rows if r["state"] != "ok")
    return _result("Employee Documents Expiry", cols, rows, None,
                   [("Documents", len(rows), "int"), ("Expired / due in %d days" % alert, due, "int")])


@report("eos_liability", GROUP_HR, "End-of-Service Liability", params=("as_of",),
        desc="If every employee left on the date: the end-of-service award (ended by the employer / if they "
             "resigned) and the unused annual leave — what the shop owes its staff.")
def r_eos_liability(date_from, date_to, **_):
    from modules import payroll as PR
    rows = []
    for x in PR.eos_liability(date_to):
        rows.append({"code": x["code"], "employee": x["name"], "join": x["join_date"], "years": x["years"],
                     "wage": x["wage"], "ended": x["award_ended"], "resigned": x["award_resigned"],
                     "leave_days": x["leave_days"], "leave_value": x["leave_value"],
                     "total": round(x["award_ended"] + x["leave_value"], 2)})
    cols = [col("code", "Code", width=70), col("employee", "Employee", width=170), col("join", "Joined", "date", 95),
            col("years", "Years", "qty", 65), col("wage", "Wage", "money", 90),
            col("ended", "Award (ended by employer)", "money", 140), col("resigned", "Award (if resigned)", "money", 125),
            col("leave_days", "Leave days", "qty", 80), col("leave_value", "Leave value", "money", 95),
            col("total", "Owed if ended today", "money", 130)]
    tot = _sum_totals(rows, ["ended", "resigned", "leave_value", "total"], "employee")
    return _result("End-of-Service Liability", cols, rows, tot,
                   [("Employees", len(rows), "int"), ("End-of-service awards", tot["ended"], "money"),
                    ("Unused leave", tot["leave_value"], "money"), ("Total owed", tot["total"], "money")],
                   ["Arts. 84–85: half a month a year for 5 years, a month a year after, on the last actual wage "
                    "(Settings ▸ HR). Keep this money aside — or book it monthly as End of Service Provision (2330)."])


# ── v7.0: country and business-module reports register themselves ──
for _m in ("modules.country_reports", "modules.restaurant_reports",
           "modules.production_reports", "modules.textile_reports"):
    try:
        __import__(_m)
    except ImportError as _e:
        if _m.split(".")[-1] not in str(_e):        # a real error inside the module
            print("[DayBill] %s not loaded: %s" % (_m, _e))
    except Exception as _e:
        print("[DayBill] %s not loaded: %s" % (_m, _e))
