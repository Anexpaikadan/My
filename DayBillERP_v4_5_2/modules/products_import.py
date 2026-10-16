# ============================================================
# DayBill ERP - Products bulk import / export template (Excel)
# Creates a template with ALL product fields, and imports rows
# into the products table (insert new / update existing by code).
# ============================================================

import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from config.database import db

# Column order in the Excel template. Header -> products field.
# Required: product_code, product_name_en, unit_price
COLUMNS = [
    ("Product Code*", "product_code"),
    ("Name (English)*", "product_name_en"),
    ("Name (Arabic)", "product_name_ar"),
    ("Description (EN)", "description_en"),
    ("Description (AR)", "description_ar"),
    ("Alias / Quick Search", "alias_name"),
    ("Product Type (product/service)", "product_type"),
    ("Unit Price*", "unit_price"),
    ("Purchase Price", "purchase_price"),
    ("Currency", "currency"),
    ("Unit of Measure (EN)", "unit_of_measure_en"),
    ("Unit of Measure (AR)", "unit_of_measure_ar"),
    ("VAT Code (S/Z/E/O)", "vat_code"),
    ("Category", "category_name"),
    ("Barcode", "barcode"),
    ("Track Inventory (1/0)", "track_inventory"),
    ("Opening Stock", "stock_quantity"),
    ("Reorder Level", "reorder_level"),
    ("Sold by Weight (1/0)", "is_weighed"),
    ("Weight Unit (kg/g)", "weight_unit"),
    ("Scale PLU", "scale_plu"),
    ("List in Sales (1/0)", "list_in_sales"),
    ("List in Purchase (1/0)", "list_in_purchase"),
    ("Active (1/0)", "is_active"),
]

SAMPLE_ROWS = [
    ["P001", "Pepsi Cola 250 ml", "بيبسي كولا", "", "", "P250", "product",
     "2.00", "1.50", "SAR", "Piece", "حبة", "S", "Food & Beverages", "",
     "1", "100", "10", "0", "kg", "", "1", "1", "1"],
    ["VEG-TOM", "Tomato", "طماطم", "", "", "tomato", "product",
     "5.00", "3.50", "SAR", "kg", "كجم", "S", "Vegetables", "",
     "1", "50", "5", "1", "kg", "2001", "1", "1", "1"],
]


def _has_column(name):
    try:
        return bool(db.execute_scalar(
            "SELECT COUNT(*) FROM sys.columns "
            "WHERE object_id=OBJECT_ID('products') AND name=?", (name,)))
    except Exception:
        return False


def write_template(path):
    """Write an Excel template with headers + 2 sample rows."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    headers = [h for h, _ in COLUMNS]
    ws.append(headers)
    hfill = PatternFill("solid", fgColor="006C35")
    for c, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hfill
    for row in SAMPLE_ROWS:
        ws.append(row)
    # widen columns a bit
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 20
    # a small notes sheet
    ws2 = wb.create_sheet("Instructions")
    notes = [
        "DayBill ERP - Product Import Template",
        "",
        "Fields marked * are required: Product Code, Name (English), Unit Price.",
        "VAT Code: S = Standard 15%, Z = Zero-rated, E = Exempt, O = Out of scope.",
        "Track Inventory / Sold by Weight / List flags / Active: use 1 or 0.",
        "Category: must match an existing category name (or leave blank).",
        "Existing products (same Product Code) will be UPDATED, new ones INSERTED.",
        "Delete the 2 sample rows before importing your own data.",
    ]
    for n in notes:
        ws2.append([n])
    ws2.column_dimensions["A"].width = 80
    wb.save(path)


def read_rows(path):
    """Read the Excel file, return list of dicts keyed by products field."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    # map header text -> field using COLUMNS (tolerant: match by header text)
    htext_to_field = {h: f for h, f in COLUMNS}
    field_by_col = []
    for h in header:
        field_by_col.append(htext_to_field.get(h))
    out = []
    for r in rows[1:]:
        if r is None:
            continue
        rec = {}
        for ci, val in enumerate(r):
            if ci < len(field_by_col) and field_by_col[ci]:
                rec[field_by_col[ci]] = val
        # skip blank lines
        if not (rec.get("product_code") or rec.get("product_name_en")):
            continue
        out.append(rec)
    return out


def _vat_id_for_code(code, vat_cats):
    code = (code or "S").strip().upper()[:1] or "S"
    for vc in vat_cats:
        if (vc.get("category_code") or "").upper().startswith(code):
            return vc["vat_category_id"]
    return vat_cats[0]["vat_category_id"] if vat_cats else None


def _cat_id_for_name(name, cats):
    if not name:
        return None
    name = str(name).strip().lower()
    for c in cats:
        if (c.get("category_name_en") or "").strip().lower() == name:
            return c["category_id"]
    return None


def import_rows(rows):
    """
    Insert/update products. Returns (inserted, updated, errors[list of str]).
    Existing products are matched by product_code.
    """
    vat_cats = db.fetchall(
        "SELECT vat_category_id, category_code FROM vat_categories WHERE is_active=1") or []
    cats = db.fetchall(
        "SELECT category_id, category_name_en FROM product_categories WHERE is_active=1") or []
    has_purchase = _has_column("purchase_price")
    has_weight = _has_column("is_weighed")
    has_alias = _has_column("alias_name")
    has_listflags = _has_column("list_in_sales")

    inserted = updated = 0
    errors = []
    try:
        from config import settings as _cfg
        locked = _cfg.flag("barcode_system_locked", True)
    except Exception:
        locked = False
    from modules.product_lookup import norm_code, barcode_owner
    by_norm = {norm_code(r["product_code"]): r["product_code"] for r in
               (db.fetchall("SELECT product_code FROM products") or [])}

    def _system_barcodes(pid, maker, uom, price, pprice):
        """Locked barcodes: a 6-digit system code on the product + its base
        unit; the maker's barcode from the sheet becomes an EXTRA barcode."""
        from modules.barcode_gen import new_unique_barcode
        unit = db.fetchone("SELECT unit_id, barcode FROM product_units WHERE product_id=? "
                           "AND ISNULL(is_active,1)=1 ORDER BY conversion_factor", (pid,))
        if not unit:
            code_ = new_unique_barcode(exclude=[maker])
            db.execute("INSERT INTO product_units (product_id, unit_name_en, conversion_factor, "
                       "selling_price, purchase_price, barcode, is_default, is_sales_unit, "
                       "is_purchase_unit, is_active) VALUES (?,?,1,?,?,?,1,1,1,1)",
                       (pid, uom or "PCs", price, pprice, code_))
            unit = db.fetchone("SELECT unit_id, barcode FROM product_units WHERE product_id=? "
                               "ORDER BY unit_id DESC", (pid,))
        elif not (unit.get("barcode") or "").strip():
            code_ = new_unique_barcode(exclude=[maker])
            db.execute("UPDATE product_units SET barcode=? WHERE unit_id=?", (code_, unit["unit_id"]))
            unit["barcode"] = code_
        db.execute("UPDATE products SET barcode=? WHERE product_id=? AND ISNULL(barcode,'')=''",
                   (unit["barcode"], pid))
        if maker and not barcode_owner(maker):
            db.execute("INSERT INTO product_unit_barcodes (product_id, unit_id, barcode, is_active) "
                       "VALUES (?,?,?,1)", (pid, unit["unit_id"], maker))

    def num(v, d=0.0):
        try:
            return float(v)
        except Exception:
            return d

    def flag(v, d=1):
        if v is None or v == "":
            return d
        s = str(v).strip().lower()
        return 1 if s in ("1", "yes", "true", "y") else 0

    for i, r in enumerate(rows, 2):  # row 2 = first data row
        code = (str(r.get("product_code")).strip()
                if r.get("product_code") is not None else "")
        name = (str(r.get("product_name_en")).strip()
                if r.get("product_name_en") is not None else "")
        if not code or not name:
            errors.append(f"Row {i}: missing code or name — skipped.")
            continue

        data = {
            "product_code": code,
            "product_name_en": name,
            "product_name_ar": (str(r.get("product_name_ar")).strip()
                                if r.get("product_name_ar") else None),
            "description_en": (str(r.get("description_en")).strip()
                               if r.get("description_en") else None),
            "description_ar": (str(r.get("description_ar")).strip()
                               if r.get("description_ar") else None),
            "product_type": (str(r.get("product_type")).strip().lower()
                             or "product") if r.get("product_type") else "product",
            "unit_price": num(r.get("unit_price")),
            "currency": (str(r.get("currency")).strip() or "SAR")
                        if r.get("currency") else "SAR",
            "unit_of_measure_en": (str(r.get("unit_of_measure_en")).strip()
                                   or "Unit") if r.get("unit_of_measure_en") else "Unit",
            "unit_of_measure_ar": (str(r.get("unit_of_measure_ar")).strip()
                                   or "وحدة") if r.get("unit_of_measure_ar") else "وحدة",
            "vat_category_id": _vat_id_for_code(r.get("vat_code"), vat_cats),
            "category_id": _cat_id_for_name(r.get("category_name"), cats),
            "barcode": (str(r.get("barcode")).strip()
                        if r.get("barcode") else None),
            "track_inventory": flag(r.get("track_inventory"), 0),
            "stock_quantity": 0.0,          # opening stock: see below (v5.26)
            "reorder_level": num(r.get("reorder_level")),
            "is_active": flag(r.get("is_active"), 1),
        }
        if has_alias:
            data["alias_name"] = (str(r.get("alias_name")).strip()
                                  if r.get("alias_name") else None)
        if has_purchase:
            data["purchase_price"] = num(r.get("purchase_price"))
        if has_listflags:
            data["list_in_sales"] = flag(r.get("list_in_sales"), 1)
            data["list_in_purchase"] = flag(r.get("list_in_purchase"), 1)
        if has_weight:
            data["is_weighed"] = flag(r.get("is_weighed"), 0)
            data["weight_unit"] = (str(r.get("weight_unit")).strip() or "kg") \
                if r.get("weight_unit") else "kg"
            data["scale_plu"] = (str(r.get("scale_plu")).strip()
                                 if r.get("scale_plu") else None)

        # the same code in another format (p001 = P1) is the same product
        code = by_norm.get(norm_code(code), code)
        by_norm[norm_code(code)] = code
        data["product_code"] = code
        maker = data.get("barcode") or ""
        if locked:
            data.pop("barcode", None)          # system barcodes are never imported
        try:
            exists = db.execute_scalar(
                "SELECT product_id FROM products WHERE product_code=?", (code,))
            try:
                from modules import price_watch as _pw
                _before = _pw.snapshot([exists] if exists else [])
            except Exception:
                _before = None
            stock_given = r.get("stock_quantity") not in (None, "")
            open_qty = num(r.get("stock_quantity"))
            if exists:
                data.pop("stock_quantity", None)   # never overwrite a product's stock
                set_clause = ", ".join(f"{k}=?" for k in data if k != "product_code")
                vals = [v for k, v in data.items() if k != "product_code"]
                db.execute(
                    f"UPDATE products SET {set_clause}, updated_at=GETDATE() "
                    f"WHERE product_code=?", (*vals, code))
                updated += 1
            else:
                cols = ", ".join(data.keys())
                ph = ", ".join(["?"] * len(data))
                db.execute(f"INSERT INTO products ({cols}) VALUES ({ph})",
                           tuple(data.values()))
                inserted += 1
            if locked:
                try:
                    pid_s = exists or db.execute_scalar(
                        "SELECT product_id FROM products WHERE product_code=?", (code,))
                    _system_barcodes(pid_s, maker, data.get("unit_of_measure_en"),
                                     data.get("unit_price") or 0, data.get("purchase_price") or 0)
                except Exception as e_:
                    errors.append(f"Row {i} ({code}) barcode: {str(e_)[:80]}")
            # Opening Stock column → opening stock (stock card + books), for
            # new products and products that have no purchases / sales yet
            if stock_given and str(data.get("product_type") or "product") != "service":
                try:
                    pid_o = exists or db.execute_scalar(
                        "SELECT product_id FROM products WHERE product_code=?", (code,))
                    from modules import opening_balances as _ob
                    from modules import report_center as _rc
                    moved = [m for m in _rc._movements("1900-01-01", None, pid_o)
                             if m.get("ref") != "opening"]
                    if not moved:
                        # nothing but the opening → stock = opening stock
                        db.execute("UPDATE products SET stock_quantity=? WHERE product_id=?",
                                   (_ob.opening_qty(pid_o), pid_o))
                        if exists or abs(open_qty) > 1e-9:
                            _ob.set_opening_stock(pid_o, open_qty,
                                                  data.get("purchase_price") or None)
                    elif abs(open_qty - _ob.opening_qty(pid_o)) > 1e-9:
                        errors.append(f"Row {i} ({code}): stock not changed — this product "
                                      f"already has purchases / sales (use Stock Adjustment).")
                except Exception as e_:
                    errors.append(f"Row {i} ({code}) opening stock: {str(e_)[:80]}")
            try:
                if _before is not None:
                    pid_ = exists or db.execute_scalar(
                        "SELECT product_id FROM products WHERE product_code=?", (code,))
                    _before[("_ids",)] = [pid_]
                    _pw.record(_before, "import", new_product=not exists)
            except Exception:
                pass
        except Exception as e:
            errors.append(f"Row {i} ({code}): {str(e)[:80]}")

    return inserted, updated, errors
